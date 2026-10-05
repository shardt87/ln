"""Build the one-page leave-behinds from one layout spec:

    pages/pdf/<KEY>_one_pager.pdf        print-ready (US Letter, vector text, 320-dpi render)
    pages/pdf/All_one_pagers.pdf         all pages, one file
    pages/pptx/<KEY>_one_pager.pptx      editable (text, callouts, legend are native shapes)
    pages/pptx/All_one_pagers.pptx       all pages in one deck (one slide = one page)
    renders/annotated/<VIEW>.png         render with the same callouts burned in (for email/screens)

    python3 pages.py [KEY ...]
"""
import os
import sys
import json

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from content import PAGES, COMMON  # noqa: E402

ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
RCLEAN = os.environ.get("SW_RENDER_DIR", os.path.join(ROOT, "renders", "clean"))
FONT_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

PW, PH = 8.5, 11.0
ML = 0.5
INK = "#1F2A33"         # body text
MUTED = "#5B6670"
ACCENT = "#B5121B"      # restrained red rule (align with approved brand palette)
MARK = "#23303B"        # callout markers
PANEL = "#F6F7F8"       # render background
RULE = "#D5D9DC"
SWATCH = {"Wire_Control": "#2E66B0", "Jacket_Power": "#1C1D1F", "Wire_Ground": "#3E8E41",
          "Highlight_Field": "#D88A1E"}
LEGEND_TXT = dict(COMMON["legend"])

# render panel (middle 50 %)
IMG_X, IMG_Y, IMG_W = ML + 0.035, 2.02, 7.43
IMG_H = IMG_W * 1664 / 2400


def hexrgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ---------------------------------------------------------------------------
# layout spec (inches, origin top-left)
# ---------------------------------------------------------------------------
def build_spec(page):
    view = page["view"]
    meta = json.load(open(os.path.join(RCLEAN, view + ".json")))
    sx, sy = meta["size"]
    el = []
    T = lambda **k: el.append(dict(kind="text", **k))
    R = lambda **k: el.append(dict(kind="rect", **k))
    # ---- top band (0 - 1.65 in) ----
    R(x=ML, y=0.38, w=1.75, h=0.48, fill=None, line="#9AA3AA", dash=True, name="Logo_Placeholder")
    T(x=ML, y=0.38, w=1.75, h=0.48, text="SOUTHWIRE LOGO\nplaceholder - insert approved artwork", size=7,
      color=MUTED, align="center", valign="middle", name="Logo_Placeholder_Text")
    T(x=2.45, y=0.40, w=3.0, h=0.22, text=COMMON["unit"].upper(), size=8.5, bold=True, color=ACCENT,
      name="Business_Unit")
    T(x=4.9, y=0.40, w=PW - ML - 4.9, h=0.22, text=f"Discussion with {page['company']}", size=8.5,
      color=MUTED, align="right", name="Discussion_Line")
    hs = 21.0
    while _textw(page["headline"], hs, True) > PW - 2 * ML - 0.05 and hs > 16:
        hs -= 0.5
    T(x=ML, y=0.98, w=PW - 2 * ML, h=0.42, text=page["headline"], size=hs, bold=True, color=INK,
      name="Headline")
    R(x=ML, y=1.52, w=0.9, h=0.035, fill=ACCENT, line=None, name="Accent_Rule")
    # ---- middle (1.65 - 7.15 in) ----
    T(x=ML, y=1.62, w=PW - 2 * ML, h=0.36, text=page["support"], size=10, color=INK, name="Supporting_Sentence")
    el.append(dict(kind="image", x=IMG_X, y=IMG_Y, w=IMG_W, h=IMG_H, path=os.path.join(RCLEAN, view + ".png"),
                   name="Render_" + view))
    # callouts
    for c in meta["callouts"]:
        n = c["n"]
        ax = IMG_X + c["anchor_px"][0] / sx * IMG_W
        ay = IMG_Y + c["anchor_px"][1] / sy * IMG_H
        lab_n = page.get("labels", {}).get(n, c["label_norm"])
        c["label_norm"] = lab_n
        mx = IMG_X + lab_n[0] * IMG_W
        my = IMG_Y + lab_n[1] * IMG_H
        el.append(dict(kind="leader", x1=mx, y1=my, x2=ax, y2=ay, name=f"Callout_{n}_Leader"))
        el.append(dict(kind="dot", x=ax, y=ay, r=0.035, name=f"Callout_{n}_Anchor"))
        el.append(dict(kind="marker", x=mx, y=my, r=0.135, n=n, name=f"Callout_{n}_Marker"))
        label = page["callouts"][n - 1]
        w = min(2.15, 0.16 + _textw(label, 8.5, True))
        lines = 1 if w < 2.15 else 2
        h = 0.24 if lines == 1 else 0.40
        left = (lab_n[2] == "L") if len(lab_n) > 2 else lab_n[0] > 0.55
        lx = mx - 0.17 - w if left else mx + 0.17
        lx = max(IMG_X + 0.05, min(lx, IMG_X + IMG_W - w - 0.05))
        R(x=lx, y=my - h / 2, w=w, h=h, fill="#FFFFFF", line=RULE, round=True, name=f"Callout_{n}_LabelBox")
        T(x=lx + 0.06, y=my - h / 2, w=w - 0.12, h=h, text=label, size=8.5, bold=True, color=INK,
          valign="middle", align="left", name=f"Callout_{n}_Label")
    # legend (inside render panel, bottom-left)
    keys = page["legend"]
    lh = 0.17 * len(keys) + 0.30
    pos = page.get("legend_pos", "bl")
    lgx = IMG_X + 0.10 if pos[1] == "l" else IMG_X + IMG_W - 2.62 - 0.04
    lgy = IMG_Y + IMG_H - 0.14 - 0.17 * len(keys) - 0.20 if pos[0] == "b" else IMG_Y + 0.11
    R(x=lgx - 0.06, y=lgy - 0.07, w=2.62, h=0.17 * len(keys) + 0.30, fill="#FFFFFF", line=RULE, name="Legend_Box",
      alpha=0.92)
    for i, k in enumerate(keys):
        R(x=lgx, y=lgy + i * 0.17 + 0.02, w=0.22, h=0.07, fill=SWATCH[k], line=None, name=f"Legend_Swatch_{i+1}")
        T(x=lgx + 0.30, y=lgy + i * 0.17 - 0.03, w=2.2, h=0.17, text=LEGEND_TXT[k], size=7, color=INK,
          name=f"Legend_Text_{i+1}")
    T(x=lgx, y=lgy + len(keys) * 0.17 - 0.01, w=2.5, h=0.24, text=COMMON["legend_note"], size=6, color=MUTED,
      name="Legend_Note")
    # ---- lower 25 % (7.15 - 9.9 in) ----
    y0 = 7.30
    T(x=ML, y=y0, w=4, h=0.2, text="APPLICATIONS TO EVALUATE", size=7.5, bold=True, color=ACCENT,
      name="Section_Label")
    R(x=ML, y=y0 + 0.24, w=PW - 2 * ML, h=0.01, fill=RULE, line=None, name="Section_Rule")
    colw = (PW - 2 * ML - 0.4) / 3
    for i in range(3):
        cx = ML + i * (colw + 0.2)
        el.append(dict(kind="marker", x=cx + 0.135, y=y0 + 0.53, r=0.135, n=i + 1, name=f"Statement_{i+1}_Marker"))
        T(x=cx + 0.34, y=y0 + 0.38, w=colw - 0.34, h=0.42, text=page["callouts"][i], size=9.5, bold=True, color=INK,
          valign="middle", name=f"Statement_{i+1}_Title")
        T(x=cx, y=y0 + 0.86, w=colw, h=0.80, text=page["statements"][i], size=9, color=INK,
          name=f"Statement_{i+1}_Text")
    T(x=ML, y=9.05, w=PW - 2 * ML, h=0.18, text=COMMON["established"], size=7.5, bold=True, color=INK,
      name="Established_Capability")
    T(x=ML, y=9.23, w=PW - 2 * ML, h=0.18, text=COMMON["proposed"], size=7.5, color=INK, name="Proposed_Scope")
    T(x=ML, y=9.45, w=PW - 2 * ML, h=0.30, text=COMMON["disclaimer"], size=6.5, color=MUTED, name="Disclaimer")
    # ---- bottom 10 % (9.9 - 11 in) ----
    R(x=ML, y=9.90, w=5.15, h=0.78, fill="#EEF1F3", line=None, name="NextStep_Box")
    R(x=ML, y=9.90, w=0.05, h=0.78, fill=ACCENT, line=None, name="NextStep_Bar")
    T(x=ML + 0.16, y=9.95, w=4.9, h=0.18, text="PROPOSED NEXT STEP", size=7.5, bold=True, color=ACCENT,
      name="NextStep_Label")
    T(x=ML + 0.16, y=10.13, w=4.9, h=0.52, text=page["next"], size=9, color=INK, name="NextStep_Text")
    cx = ML + 5.35
    T(x=cx, y=9.93, w=PW - ML - cx, h=0.2, text=COMMON["contact_name"], size=9.5, bold=True, color=INK,
      name="Contact_Name")
    T(x=cx, y=10.13, w=PW - ML - cx, h=0.32, text=COMMON["contact_title"], size=7.5, color=INK, name="Contact_Title")
    T(x=cx, y=10.43, w=PW - ML - cx, h=0.18, text=f"{COMMON['contact_phone']}   {COMMON['contact_email']}",
      size=7.5, color=MUTED, name="Contact_Details")
    return el, meta


MODULE_LABEL = {
    "IEM": "Switchgear / control cabinet + wire kit (Module 1)",
    "EPD": "Switchgear / control cabinet + wire kit (Module 1)",
    "PATRIOT": "Control compartment terminations (Module 1)",
    "MAVERICK": "Switchgear lineup, cable entry (Module 1)",
    "CAT": "Generator package interfaces (Module 2)",
    "TAYLOR": "Generator package + kits (Module 2)",
    "POWELL": "E-house, exploded (Module 3)",
    "SIEMENS": "E-house / skid boundaries (Module 3)",
    "NVENT": "E-house, roof removed (Module 3)",
    "WESCO": "E-house modules + staged supply (Module 3)",
    "RESA": "Transformer interface, Configuration B (Module 4)",
    "ASCO": "Transfer equipment, Configuration A",
    "INPOWER": "Docking station, Configuration D (Module 5)",
    "MOSEBACH": "Load-bank testing, Configuration D (Module 5)",
    "WINAR": "Cable assembly bench (Module 6)",
}


def thumb(view):
    from PIL import Image
    out = os.path.join(ROOT, "renders", "thumbs", view + ".jpg")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    src = os.path.join(RCLEAN, view + ".png")
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(src):
        im = Image.open(src).convert("RGB")
        im.thumbnail((420, 420))
        im.save(out, quality=88)
    return out


def index_spec(pages):
    """Index page placed first in the combined PDF / PPTX."""
    el = []
    T = lambda **k: el.append(dict(kind="text", **k))
    R = lambda **k: el.append(dict(kind="rect", **k))
    R(x=ML, y=0.38, w=1.75, h=0.48, fill=None, line="#9AA3AA", dash=True, name="Logo_Placeholder")
    T(x=ML, y=0.38, w=1.75, h=0.48, text="SOUTHWIRE LOGO\nplaceholder - insert approved artwork", size=7,
      color=MUTED, align="center", valign="middle", name="Logo_Placeholder_Text")
    T(x=2.45, y=0.40, w=3.0, h=0.22, text=COMMON["unit"].upper(), size=8.5, bold=True, color=ACCENT,
      name="Business_Unit")
    T(x=ML, y=0.98, w=PW - 2 * ML, h=0.42, text="Application one-pagers: index", size=21, bold=True,
      color=INK, name="Headline")
    R(x=ML, y=1.52, w=0.9, h=0.035, fill=ACCENT, line=None, name="Accent_Rule")
    T(x=ML, y=1.62, w=PW - 2 * ML, h=0.36,
      text=(f"{len(pages)} discussion pages built on one reusable 3D model. Each page pairs one equipment view "
            "with three applications to evaluate and a proposed next step."),
      size=10, color=INK, name="Intro")
    y0 = 2.18
    for txt, x, w, al in (("#", ML, 0.3, "left"), ("COMPANY / MODEL VIEW", 1.68, 2.3, "left"),
                          ("APPLICATION", 4.08, 3.0, "left"), ("PAGE", PW - ML - 0.5, 0.5, "right")):
        T(x=x, y=y0, w=w, h=0.18, text=txt, size=7.5, bold=True, color=ACCENT, align=al, name="Header_" + txt[:4])
    R(x=ML, y=y0 + 0.22, w=PW - 2 * ML, h=0.01, fill=RULE, line=None, name="Header_Rule")
    rh = 0.535
    for i, p in enumerate(pages):
        y = y0 + 0.30 + i * rh
        el.append(dict(kind="marker", x=ML + 0.13, y=y + rh / 2 - 0.03, r=0.12, n=i + 1, name=f"Row_{i+1}_Number"))
        el.append(dict(kind="image", x=ML + 0.36, y=y + 0.01, w=0.70, h=0.70 * 1664 / 2400,
                       path=thumb(p["view"]), name=f"Row_{i+1}_Thumb"))
        T(x=1.68, y=y + 0.04, w=2.3, h=0.2, text=p["company"], size=9.5, bold=True, color=INK,
          name=f"Row_{i+1}_Company")
        T(x=1.68, y=y + 0.24, w=2.3, h=0.22, text=MODULE_LABEL.get(p["key"], ""), size=7, color=MUTED,
          name=f"Row_{i+1}_View")
        T(x=4.08, y=y + 0.04, w=3.0, h=0.42, text=p["headline"], size=8.5, color=INK, name=f"Row_{i+1}_Application")
        T(x=PW - ML - 0.5, y=y + 0.12, w=0.5, h=0.2, text=str(i + 2), size=10, bold=True, color=INK,
          align="right", name=f"Row_{i+1}_Page")
        if i < len(pages) - 1:
            R(x=ML, y=y + rh - 0.035, w=PW - 2 * ML, h=0.006, fill="#E6E9EB", line=None, name=f"Row_{i+1}_Rule")
    T(x=ML, y=10.42, w=PW - 2 * ML, h=0.34,
      text=("Discussion concepts only. Company names identify the intended conversation; no partnership, pilot, "
            "purchase or product qualification is implied. " + COMMON["disclaimer"]),
      size=6.5, color=MUTED, name="Index_Disclaimer")
    return el, None, {"key": "INDEX", "company": "Index", "view": "-", "headline": "Index"}


def numbered(spec, n, total):
    el, meta, page = spec
    el = el + [dict(kind="text", x=PW - ML - 1.2, y=10.74, w=1.2, h=0.16, text=f"Page {n} of {total}", size=6.5,
                    color=MUTED, align="right", name="Page_Number")]
    return el, meta, page


_FONTS = {}


def _pil_font(size_px, bold):
    from PIL import ImageFont
    k = (size_px, bold)
    if k not in _FONTS:
        _FONTS[k] = ImageFont.truetype(FONT_B if bold else FONT_R, size_px)
    return _FONTS[k]


def _textw(s, pt, bold):
    f = _pil_font(200, bold)
    return f.getlength(s) / 200 * pt / 72.0


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------
def to_pdf(specs, path):
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor, Color
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, Frame
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    if "LS" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("LS", FONT_R))
        pdfmetrics.registerFont(TTFont("LS-B", FONT_B))
    c = canvas.Canvas(path, pagesize=(PW * 72, PH * 72))
    c.setTitle("Southwire Power Generation Solutions - application one-pagers")
    c.setAuthor("Southwire Power Generation Solutions")
    I = lambda v: v * 72
    Y = lambda y: (PH - y) * 72
    overflow = []
    for el, meta, page in specs:
        c.setFillColor(HexColor("#FFFFFF"))
        c.rect(0, 0, I(PW), I(PH), stroke=0, fill=1)
        for e in el:
            k = e["kind"]
            if k == "image":
                c.drawImage(e["path"], I(e["x"]), Y(e["y"] + e["h"]), I(e["w"]), I(e["h"]))
            elif k == "rect":
                if e.get("fill"):
                    col = HexColor(e["fill"])
                    if e.get("alpha"):
                        col = Color(col.red, col.green, col.blue, alpha=e["alpha"])
                    c.setFillColor(col)
                if e.get("line"):
                    c.setStrokeColor(HexColor(e["line"]))
                    c.setLineWidth(0.6)
                    c.setDash(3, 2) if e.get("dash") else c.setDash()
                args = (I(e["x"]), Y(e["y"] + e["h"]), I(e["w"]), I(e["h"]))
                if e.get("round"):
                    c.roundRect(*args, 4, stroke=1 if e.get("line") else 0, fill=1 if e.get("fill") else 0)
                else:
                    c.rect(*args, stroke=1 if e.get("line") else 0, fill=1 if e.get("fill") else 0)
                c.setDash()
            elif k == "leader":
                c.setStrokeColor(HexColor(MARK))
                c.setLineWidth(1.0)
                c.line(I(e["x1"]), Y(e["y1"]), I(e["x2"]), Y(e["y2"]))
            elif k == "dot":
                c.setFillColor(HexColor(MARK))
                c.setStrokeColor(HexColor("#FFFFFF"))
                c.setLineWidth(1.0)
                c.circle(I(e["x"]), Y(e["y"]), I(e["r"]), stroke=1, fill=1)
            elif k == "marker":
                c.setFillColor(HexColor(MARK))
                c.setStrokeColor(HexColor("#FFFFFF"))
                c.setLineWidth(1.2)
                c.circle(I(e["x"]), Y(e["y"]), I(e["r"]), stroke=1, fill=1)
                c.setFillColor(HexColor("#FFFFFF"))
                c.setFont("LS-B", 10.5)
                c.drawCentredString(I(e["x"]), Y(e["y"]) - 3.7, str(e["n"]))
            elif k == "text":
                al = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT}[e.get("align", "left")]
                st = ParagraphStyle("s", fontName="LS-B" if e.get("bold") else "LS", fontSize=e["size"],
                                    leading=e["size"] * 1.22, textColor=HexColor(e.get("color", INK)), alignment=al)
                txt = e["text"].replace("&", "&amp;").replace("<", "&lt;").replace("\n", "<br/>")
                p = Paragraph(txt, st)
                w, h = p.wrap(I(e["w"]), I(e["h"]) * 4)
                if h > I(e["h"]) + 1.5:
                    overflow.append((page["key"], e["name"], round(h / 72, 3), e["h"]))
                yy = e["y"]
                if e.get("valign") == "middle":
                    yy = e["y"] + (e["h"] - h / 72) / 2
                p.drawOn(c, I(e["x"]), Y(yy) - h)
        c.showPage()
    c.save()
    return overflow


# ---------------------------------------------------------------------------
# PPTX (python-pptx) -- every element native and editable
# ---------------------------------------------------------------------------
def to_pptx(specs, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PW), Inches(PH)
    blank = prs.slide_layouts[6]
    RGB = lambda h: RGBColor(*hexrgb(h))
    for el, meta, page in specs:
        s = prs.slides.add_slide(blank)
        for e in el:
            k = e["kind"]
            if k == "image":
                pic = s.shapes.add_picture(e["path"], Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                pic.name = e["name"]
            elif k == "rect":
                shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e.get("round") else MSO_SHAPE.RECTANGLE,
                                         Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                if e.get("round"):
                    shp.adjustments[0] = 0.18
                if e.get("fill"):
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = RGB(e["fill"])
                    if e.get("alpha"):
                        sf = shp.fill._xPr.find(qn("a:solidFill"))
                        clr = sf[0]
                        a = clr.makeelement(qn("a:alpha"), {"val": str(int(e["alpha"] * 100000))})
                        clr.append(a)
                else:
                    shp.fill.background()
                if e.get("line"):
                    shp.line.color.rgb = RGB(e["line"])
                    shp.line.width = Pt(0.6)
                    if e.get("dash"):
                        from pptx.enum.dml import MSO_LINE
                        shp.line.dash_style = MSO_LINE.DASH
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                shp.name = e["name"]
            elif k == "leader":
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x1"]), Inches(e["y1"]),
                                            Inches(e["x2"]), Inches(e["y2"]))
                ln.line.color.rgb = RGB(MARK)
                ln.line.width = Pt(1.0)
                ln.name = e["name"]
            elif k in ("dot", "marker"):
                r = e["r"]
                shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(e["x"] - r), Inches(e["y"] - r), Inches(2 * r), Inches(2 * r))
                shp.fill.solid()
                shp.fill.fore_color.rgb = RGB(MARK)
                shp.line.color.rgb = RGB("#FFFFFF")
                shp.line.width = Pt(1.2 if k == "marker" else 1.0)
                shp.shadow.inherit = False
                shp.name = e["name"]
                if k == "marker":
                    tf = shp.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = str(e["n"])
                    run.font.size = Pt(10.5)
                    run.font.bold = True
                    run.font.name = "Arial"
                    run.font.color.rgb = RGB("#FFFFFF")
            elif k == "text":
                tb = s.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
                tb.name = e["name"]
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE if e.get("valign") == "middle" else MSO_ANCHOR.TOP
                for i, line in enumerate(e["text"].split("\n")):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                   "right": PP_ALIGN.RIGHT}[e.get("align", "left")]
                    p.line_spacing = 1.08
                    run = p.add_run()
                    run.text = line
                    run.font.size = Pt(e["size"])
                    run.font.bold = bool(e.get("bold"))
                    run.font.name = "Arial"
                    run.font.color.rgb = RGB(e.get("color", INK))
        if page["key"] == "INDEX":
            s.notes_slide.notes_text_frame.text = "Index of the company one-pagers; page numbers refer to this file."
            continue
        s.notes_slide.notes_text_frame.text = (
            f"View: {page['view']} (master .blend). Discussion concept for {page['company']}; "
            "no agreement, pilot or purchase is implied. Replace the logo placeholder and the [Phone]/[Email] "
            "placeholders before use.")
    prs.save(path)


# ---------------------------------------------------------------------------
# Annotated render PNG (same callouts, burned in)
# ---------------------------------------------------------------------------
def annotate(page, meta):
    from PIL import Image, ImageDraw
    view = page["view"]
    im = Image.open(os.path.join(RCLEAN, view + ".png")).convert("RGB")
    sx, sy = meta["size"]
    ppi = sx / IMG_W
    d = ImageDraw.Draw(im)
    for c in meta["callouts"]:
        c["label_norm"] = page.get("labels", {}).get(c["n"], c["label_norm"])
        ax, ay = c["anchor_px"]
        mx, my = c["label_norm"][0] * sx, c["label_norm"][1] * sy
        d.line([(mx, my), (ax, ay)], fill=hexrgb(MARK), width=int(ppi / 72 * 1.0) + 1)
        r = 0.035 * ppi
        d.ellipse([ax - r, ay - r, ax + r, ay + r], fill=hexrgb(MARK), outline=(255, 255, 255), width=3)
        R = 0.135 * ppi
        d.ellipse([mx - R, my - R, mx + R, my + R], fill=hexrgb(MARK), outline=(255, 255, 255), width=5)
        f = _pil_font(int(10.5 / 72 * ppi), True)
        d.text((mx, my), str(c["n"]), font=f, fill=(255, 255, 255), anchor="mm")
        label = page["callouts"][c["n"] - 1]
        fl = _pil_font(int(8.5 / 72 * ppi), True)
        w = min(2.15 * ppi, 0.16 * ppi + d.textlength(label, font=fl))
        two = w >= 2.15 * ppi
        h = (0.40 if two else 0.24) * ppi
        ln_ = c["label_norm"]
        left = (ln_[2] == "L") if len(ln_) > 2 else ln_[0] > 0.55
        lx = mx - 0.17 * ppi - w if left else mx + 0.17 * ppi
        lx = max(0.05 * ppi, min(lx, sx - w - 0.05 * ppi))
        d.rounded_rectangle([lx, my - h / 2, lx + w, my + h / 2], radius=int(0.05 * ppi), fill=(255, 255, 255),
                            outline=hexrgb(RULE), width=3)
        if two:
            words, lines, cur = label.split(), [], ""
            for wd in words:
                t = (cur + " " + wd).strip()
                if d.textlength(t, font=fl) > w - 0.12 * ppi and cur:
                    lines.append(cur)
                    cur = wd
                else:
                    cur = t
            lines.append(cur)
            for i, ln in enumerate(lines[:2]):
                d.text((lx + 0.06 * ppi, my + (i - 0.5) * 0.13 * ppi), ln, font=fl, fill=hexrgb(INK), anchor="lm")
        else:
            d.text((lx + 0.06 * ppi, my), label, font=fl, fill=hexrgb(INK), anchor="lm")
    out = os.path.join(ROOT, "renders", "annotated", f"{view}.png")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    im.save(out, dpi=(320, 320))


def main(keys=None):
    pages = [p for p in PAGES if not keys or p["key"] in keys]
    specs = []
    for p in pages:
        el, meta = build_spec(p)
        specs.append((el, meta, p))
        annotate(p, meta)
    os.makedirs(os.path.join(ROOT, "pages", "pdf"), exist_ok=True)
    os.makedirs(os.path.join(ROOT, "pages", "pptx"), exist_ok=True)
    over = []
    for sp in specs:
        k = sp[2]["key"]
        over += to_pdf([sp], os.path.join(ROOT, "pages", "pdf", f"{k}_one_pager.pdf"))
        to_pptx([sp], os.path.join(ROOT, "pages", "pptx", f"{k}_one_pager.pptx"))
    if not keys:
        total = len(specs) + 1
        allspecs = [numbered(index_spec(pages), 1, total)] + [numbered(sp, i + 2, total) for i, sp in enumerate(specs)]
        over += to_pdf(allspecs, os.path.join(ROOT, "pages", "pdf", "All_one_pagers.pdf"))
        to_pptx(allspecs, os.path.join(ROOT, "pages", "pptx", "All_one_pagers.pptx"))
    for o in over:
        print("OVERFLOW", o)
    print("pages:", [s[2]["key"] for s in specs])


if __name__ == "__main__":
    main(sys.argv[1:] or None)
