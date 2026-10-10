"""Cable schedule and cable BOM for the SK-3X1 campus, derived from the 3D model (conceptual, not engineered).

1. Load list: every powered item of the model (sk3x1_model.json, data-centre option excluded) is matched to an
   equipment class that gives its power circuits (voltage, typical rating, number of circuits, source bus) and its
   control / instrumentation / fire-alarm / fibre / thermocouple cables. Ratings are typical for the class unless
   the drawing names one (e.g. "BF-A: booster fan (~16 MW)").
2. Routing: the electrical routes of the model (cable trays, duct banks, buried and surface cable, trenches) form a
   graph; each cable takes the shortest path from its source to its load along that graph, plus the drops from the
   tray / bank to the equipment terminals. Run length = (route + vertical + terminations) x (1 + slack).
3. Sizing: conductor size from ampacity (125 % of full-load current for motors and feeders) and, for 480 V, a 3 %
   voltage-drop check; parallel sets above the largest practical size.
4. Products: typical Southwire product families for each application (ARMOR-X MC-HL in Class I Div 2 and turbine
   areas, MV-105 elsewhere, Type TC-ER, PLTC / ITC, FPLR, Type SHD-GC / W on the portable pad, HV XLPE at 230 kV).

    python3 cable_schedule.py        -> brochure/SK-3X1_Cable_Schedule_and_BOM.xlsx
"""
import heapq
import json
import math
import os
import re
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
M = json.load(open(os.path.join(HERE, "sk3x1_model.json")))
ITEMS = [it for it in M["items"] if it.get("area") != "L" and not it["layer"].startswith("OPT_DC")]
BY_TAG = {it["tag"]: it for it in ITEMS if it.get("tag")}
AREA = M["areas"]


def find(pat):
    r = re.compile(pat)
    return [it for it in ITEMS if r.search(it["name"])]


def one(pat):
    f = find(pat)
    return f[0] if f else None


def centre(it):
    f = it["fp"]
    return ((f[0] + f[1]) / 2, (f[2] + f[3]) / 2)


# ------------------------------------------------------------------------------------------------ route graph
ELEC = {"mv_tray": "Tray", "lv_tray": "Tray", "control_tray": "Tray", "duct_bank": "Duct bank",
        "mvlv_cable": "Duct bank", "hv_cable": "HV duct bank", "cable_trench": "Trench"}


class Graph:
    def __init__(self, routes):
        self.adj = defaultdict(list)
        self.kind = {}
        segs = []
        for r in routes:
            if r["type"] not in ELEC:
                continue
            z = r["z"]
            kind = "Surface" if r.get("surface") else ELEC[r["type"]]
            P = [(round(p[0], 1), round(p[1], 1), z) for p in r["points"]]
            for a, b in zip(P, P[1:]):
                segs.append((a, b, kind))
        # split each segment at every route vertex lying on it (T-junctions), then join coincident / stacked nodes
        verts = {s[0] for s in segs} | {s[1] for s in segs}
        grid = defaultdict(list)
        for v in verts:
            grid[(int(v[0] // 20), int(v[1] // 20))].append(v)

        def near(x, y, rad):
            gx, gy = int(x // 20), int(y // 20)
            for i in range(gx - 1, gx + 2):
                for j in range(gy - 1, gy + 2):
                    for v in grid[(i, j)]:
                        if abs(v[0] - x) <= rad and abs(v[1] - y) <= rad:
                            yield v
        for a, b, kind in segs:
            horiz = abs(a[1] - b[1]) < abs(a[0] - b[0])
            lo, hi = (min(a[0], b[0]), max(a[0], b[0])) if horiz else (min(a[1], b[1]), max(a[1], b[1]))
            pts = [a, b]
            n = max(1, int((hi - lo) // 20))
            for k in range(n + 1):                      # sample along the segment for vertices on it
                t = k / n
                x, y = a[0] + (b[0] - a[0]) * t, a[1] + (b[1] - a[1]) * t
                for v in near(x, y, 12):
                    if (lo - .5 <= (v[0] if horiz else v[1]) <= hi + .5 and
                            abs((v[1] - a[1]) if horiz else (v[0] - a[0])) <= 3):
                        q = (v[0], a[1], a[2]) if horiz else (a[0], v[1], a[2])
                        pts.append(q)
            pts = sorted(set(pts), key=lambda p: (p[0], p[1]) if horiz else (p[1], p[0]))
            for p, q in zip(pts, pts[1:]):
                d = math.dist(p[:2], q[:2])
                self.adj[p].append((q, d))
                self.adj[q].append((p, d))
                self.kind[(p, q)] = self.kind[(q, p)] = kind
        nodes = list(self.adj)
        ng = defaultdict(list)
        for v in nodes:
            ng[(int(v[0] // 6), int(v[1] // 6))].append(v)
        for v in nodes:                                  # vertical / stacked transitions (risers, tray to bank)
            gx, gy = int(v[0] // 6), int(v[1] // 6)
            for i in range(gx - 1, gx + 2):
                for j in range(gy - 1, gy + 2):
                    for w in ng[(i, j)]:
                        if w != v and math.dist(v[:2], w[:2]) <= 4:
                            d = math.dist(v[:2], w[:2]) + abs(v[2] - w[2])
                            self.adj[v].append((w, d))
                            self.kind[(v, w)] = "Riser"
        # bridge separate pieces of the network (tray ends to a nearby bank or tray) with short conduit links
        comp = {}
        for v in nodes:
            if v in comp:
                continue
            stack, comp[v] = [v], v
            while stack:
                u = stack.pop()
                for w, _ in self.adj[u]:
                    if w not in comp:
                        comp[w] = v
                        stack.append(w)
        bg = defaultdict(list)
        for v in nodes:
            bg[(int(v[0] // 110), int(v[1] // 110))].append(v)
        best = {}
        for v in nodes:
            gx, gy = int(v[0] // 110), int(v[1] // 110)
            for i in range(gx - 1, gx + 2):
                for j in range(gy - 1, gy + 2):
                    for w in bg[(i, j)]:
                        if comp[w] != comp[v]:
                            d = abs(v[0] - w[0]) + abs(v[1] - w[1])
                            key = tuple(sorted((comp[v], comp[w])))
                            if d <= 110 and (key not in best or d < best[key][0]):
                                best[key] = (d, v, w)
        for d, v, w in best.values():
            c = d * 1.2 + abs(v[2] - w[2])
            self.adj[v].append((w, c))
            self.adj[w].append((v, c))
            self.kind[(v, w)] = self.kind[(w, v)] = "Conduit link"
        self.n_links = len(best)
        self.nodes = nodes
        ag = defaultdict(list)
        for v in nodes:
            ag[(int(v[0] // 50), int(v[1] // 50))].append(v)
        self.ag = ag

    def access(self, it, k=4):
        """The k network nodes nearest to an item's footprint, with the drop length (horizontal + vertical)."""
        f = it["fp"]
        cx, cy = centre(it)
        cand = []
        for rad in (1, 2, 4, 8, 16):
            gx, gy = int(cx // 50), int(cy // 50)
            cand = [v for i in range(gx - rad, gx + rad + 1) for j in range(gy - rad, gy + rad + 1) for v in self.ag[(i, j)]]
            if len(cand) >= k:
                break
        out = []
        for v in cand:
            dx = max(f[0] - v[0], 0, v[0] - f[1])
            dy = max(f[2] - v[1], 0, v[1] - f[3])
            out.append((dx + dy + abs(v[2] - 4) + 10, v))
        return sorted(out)[:k]

    def dijkstra(self, starts):
        dist = {}
        pq = [(d, v) for d, v in starts]
        heapq.heapify(pq)
        prev = {}
        while pq:
            d, v = heapq.heappop(pq)
            if v in dist:
                continue
            dist[v] = d
            for w, c in self.adj[v]:
                if w not in dist:
                    nd = d + c
                    heapq.heappush(pq, (nd, w))
                    prev.setdefault(w, (nd, v))
                    if nd < prev[w][0]:
                        prev[w] = (nd, v)
        return dist, prev


G = Graph(M["routes"])
_DCACHE = {}
PATHS = {}                                            # (src id, dst id) -> routed polyline (ft), for the viewer X-ray


def term(it):
    c = centre(it)
    z0, z1 = it["z"]
    return (round(c[0], 1), round(c[1], 1), round(z0 + min(6, (z1 - z0) / 2), 1))


def route_len(src, dst):
    """(route length, vertical + drops, dominant route kind) from src item to dst item along the network."""
    if src["id"] not in _DCACHE:
        acc = G.access(src)
        _DCACHE[src["id"]] = (G.dijkstra([(d, v) for d, v in acc]), acc)
    (dist, prev), acc_s = _DCACHE[src["id"]]
    best = None
    for d0, v in G.access(dst):
        if v in dist and (best is None or dist[v] + d0 < best[0]):
            best = (dist[v] + d0, v, d0)
    a, b = centre(src), centre(dst)
    if best is None:                                     # not on the network: L-shaped estimate
        PATHS[(src["id"], dst["id"])] = [term(src), (b[0], a[1], 0.0), term(dst)]
        return abs(a[0] - b[0]) + abs(a[1] - b[1]), 40, "Estimated"
    total, v, d_end = best
    kinds = defaultdict(float)
    vert = 0.0
    cur = v
    n = 0
    chain = [v]
    while cur in prev and n < 5000:
        d, p = prev[cur]
        chain.append(p)
        k = G.kind.get((p, cur), "Tray")
        seg = math.dist(p[:2], cur[:2])
        if k == "Riser":
            vert += abs(p[2] - cur[2])
        kinds[k] += seg
        cur = p
        n += 1
    PATHS[(src["id"], dst["id"])] = [term(src)] + [tuple(round(q, 1) for q in c) for c in reversed(chain)] + [term(dst)]
    d_start = min(d for d, w in acc_s)
    route = max(0.0, total - d_end - d_start - vert)
    kind = max(kinds, key=kinds.get) if kinds else "Tray"
    return route, vert + d_end + d_start, kind


# ------------------------------------------------------------------------------------------------ sizing data
SIZES_LV = ["14", "12", "10", "8", "6", "4", "2", "1", "1/0", "2/0", "3/0", "4/0", "250", "350", "500", "750"]
AMP_LV = dict(zip(SIZES_LV, [15, 20, 30, 50, 65, 85, 115, 130, 150, 175, 200, 230, 255, 310, 380, 475]))   # NEC 310.16 75 C Cu
R_LV = dict(zip(SIZES_LV, [3.07, 1.93, 1.21, .764, .491, .308, .194, .154, .122, .0967, .0766, .0608, .0515, .0367,
                           .0258, .0171]))                                                                  # ohm / kft
SIZES_MV = ["2", "1/0", "2/0", "4/0", "250", "350", "500", "750", "1000"]
AMP_MV = dict(zip(SIZES_MV, [165, 215, 245, 320, 355, 435, 525, 650, 750]))     # 15 kV MV-105 Cu triplexed (typ.)
AMP_MV35 = dict(zip(SIZES_MV, [130, 170, 195, 255, 285, 350, 420, 520, 600]))   # 35 kV MV-105 Al 1/C (typ.)
AWG = lambda s: s if "/" in s or len(s) < 3 else s + " kcmil"


def size_lv(I, L, V=480, maxsize="500"):
    for sets in (1, 2, 3, 4, 6, 8, 10, 12):
        for s in SIZES_LV[2:]:
            if AMP_LV[s] * sets >= I * 1.25:
                vd = 1.732 * I * (L / 1000) * R_LV[s] / sets / V * 100
                if vd <= 3.0:
                    return s, sets, vd
            if s == maxsize:
                break
    return maxsize, 12, 0


def size_mv(I, table=AMP_MV, derate=.85):
    for sets in (1, 2, 3, 4, 5, 6, 8, 10):
        for s in SIZES_MV:
            if table[s] * derate * sets >= I * 1.25:
                return s, sets
            if s == "1000":
                break
    return "1000", 10


# ------------------------------------------------------------------------------------------------ products
P = dict(
    MV15="MV-105 15 kV 1/C Cu, EPR 133%, Cu tape shield, PVC jacket (triplexed)",
    MV15AX="ARMOR-X MC-HL / MV-105 15 kV 3/C Cu, EPR 133%, tape shield, gnd, CWA, red PVC",
    MV5="MV-105 5 kV 3/C Cu, EPR 133%, tape shield, gnd, PVC jacket",
    MV5S="MV-105 5 kV 1/C Cu, EPR 133%, Cu tape shield, PVC jacket (triplexed)",
    MV25="MV-105 25 kV 1/C Cu, EPR 133%, Cu tape shield, PVC jacket",
    MV35="MV-105 35 kV 1/C Al, EPR 133%, 1/3 concentric neutral, LLDPE jacket",
    HV230="HV XLPE 230 kV 1/C 2500 kcmil Cu, Cu wire shield, Al laminate, HDPE jacket",
    LV="Type TC-ER 600 V 3/C Cu XHHW-2 + ground, PVC jacket",
    LVAX="ARMOR-X MC-HL 600 V 3/C Cu XHHW-2 + ground, CWA, black PVC (Class I Div 2)",
    VFD="VFD cable 600 V 3/C Cu XHHW-2, 3 symmetrical grounds, overall Cu tape shield, TC-ER",
    LV1C="XHHW-2 600 V 1/C Cu (feeder sets in tray / duct)",
    CT="Control cable 600 V Type TC-ER, Cu #14 AWG, PVC jacket",
    CTAX="ARMOR-X MC-HL 600 V control, Cu #14 AWG (Class I Div 2)",
    CTCT="Control cable 600 V Type TC-ER, Cu #10 AWG (CT / VT circuits)",
    PLTC="Instrumentation PLTC / TC-ER 300 V, #18 twisted shielded pairs, overall shield",
    ITC="Instrumentation ITC / PLTC, IS, #18 shielded pairs, blue jacket",
    TCX="Thermocouple extension KX, #16 shielded pairs, PLTC",
    RTD="RTD cable, #18 shielded triads, PLTC",
    FA="Fire alarm FPLR 2/C #14 shielded, red jacket (NEC 760)",
    FO="Fibre optic, single-mode, 12-fibre, armored, orange jacket",
    PV="PV / RHW-2 2 kV 1/C Cu, XLPE (DC)",
    SHD="Type SHD-GC 15 kV 3/C Cu, 2 grounds + GC, yellow CPE (portable)",
    W="Type W 2 kV 1/C Cu, EPDM / CPE (portable, cam-lock)",
    THHN="SIMpull THHN/THWN-2 600 V 1/C Cu (lighting, receptacles, small power)",
    BARE="Bare Cu 4/0 AWG, 19 strand (ground grid and equipment grounds)",
    ACSR="Overhead conductor ACSR 1590 kcmil 'Falcon' (typ.)",
)
CLASS = dict(MV15="MV 15 kV", MV15AX="MV 15 kV", MV5="MV 5 kV", MV5S="MV 5 kV", MV25="MV 25 kV", MV35="MV 35 kV", HV230="HV 230 kV",
             LV="LV 600 V", LVAX="LV 600 V", VFD="LV 600 V", LV1C="LV 600 V", CT="Control", CTAX="Control",
             CTCT="Control", PLTC="Instrumentation", ITC="Instrumentation", TCX="Instrumentation", RTD="Instrumentation",
             FA="Fire alarm", FO="Fibre / data", PV="DC 2 kV", SHD="Portable MV", W="Portable LV", THHN="Building wire",
             BARE="Grounding", ACSR="Overhead")
REEL = dict(MV15=2500, MV15AX=2000, MV5=2500, MV5S=2500, MV25=2500, MV35=2500, HV230=1800, LV=2500, LVAX=2000, VFD=2500,
            LV1C=2500, CT=5000, CTAX=3000, CTCT=5000, PLTC=5000, ITC=5000, TCX=5000, RTD=5000, FA=5000, FO=10000,
            PV=2500, SHD=1000, W=1000, THHN=5000, BARE=5000, ACSR=7000)

HAZ = re.compile(r"gas yard|metering|regulation|filter-separator|performance heater|M&R|ESD|fuel-gas|ULSD|LNG|H2 |"
                 r"electroly|gas conditioning|^GT\d|GFM|GT\d aux|SOS-|fuel|blending|compressor", re.I)

# ------------------------------------------------------------------------------------------------ circuits
ROWS = []


def src_of(pat):
    it = one(pat)
    assert it, pat
    return it


def add(src, dst, service, prod, cond, size, sets=1, cables_per_set=1, kw=None, volt=None, amps=None, note="",
        length=None, desc=None):
    if length is None:
        route, vert, kind = route_len(src, dst)
    else:
        route, vert, kind = length
    ROWS.append(dict(src=src, dst=dst, service=service, prod=prod, cond=cond, size=size, sets=sets,
                     cps=cables_per_set, kw=kw, volt=volt, amps=amps, route=round(route), vert=round(vert), kind=kind,
                     note=note, desc=desc, path=PATHS.get((src["id"], dst["id"]))))


def mv_feeder(src, dst, kw, kv, service, armored=False, motor=True, n=1, note=""):
    I = kw / (1.732 * kv * (.9 if motor else 1.0))
    for _ in range(n):
        if kv >= 30:
            s, sets = size_mv(I, AMP_MV35)
            add(src, dst, service, "MV35", "1/C", s, sets, 3, kw, kv * 1000, round(I), note)
        elif kv > 20:
            s, sets = size_mv(I)
            add(src, dst, service, "MV25", "1/C", s, sets, 3, kw, kv * 1000, round(I), note)
        elif kv > 6:
            s, sets = size_mv(I)
            if armored and s in ("2", "1/0", "2/0", "4/0", "250", "350", "500"):
                add(src, dst, service, "MV15AX", "3/C", s, sets, 1, kw, kv * 1000, round(I), note)
            else:
                add(src, dst, service, "MV15", "1/C", s, sets, 3, kw, kv * 1000, round(I), note)
        else:
            s, sets = size_mv(I)
            if s in ("750", "1000") or sets > 1:
                add(src, dst, service, "MV5S", "1/C", s, sets, 3, kw, kv * 1000, round(I), note)
            else:
                add(src, dst, service, "MV5", "3/C", s, sets, 1, kw, kv * 1000, round(I), note)


def lv_feeder(src, dst, kw, service, prod=None, n=1, note="", motor=True):
    I = kw / (1.732 * .48 * (.85 if motor else .95))
    route, vert, kind = route_len(src, dst)
    L = (route + vert) * 1.1
    s, sets, vd = size_lv(I, L)
    pr = prod or ("LVAX" if HAZ.search(dst["name"]) else "LV")
    for _ in range(n):
        add(src, dst, service, pr, "3/C+G", s, sets, 1, kw, 480, round(I, 1), note or (f"VD {vd:.1f} %" if vd else ""),
            length=(route, vert, kind))


def ctrl(src, dst, n7=1, n12=0, nct=0, haz=None):
    haz = HAZ.search(dst["name"]) if haz is None else haz
    for _ in range(n7):
        add(src, dst, "Control", "CTAX" if haz else "CT", "7/C", "14")
    for _ in range(n12):
        add(src, dst, "Control", "CTAX" if haz else "CT", "12/C", "14")
    for _ in range(nct):
        add(src, dst, "CT / VT", "CTCT", "4/C", "10")


def inst(src, dst, n4=1, n12=0, ntcx=0, nrtd=0, haz=None):
    haz = HAZ.search(dst["name"]) if haz is None else haz
    for _ in range(n4):
        add(src, dst, "Instrumentation", "ITC" if haz else "PLTC", "4 pr", "18")
    for _ in range(n12):
        add(src, dst, "Instrumentation", "ITC" if haz else "PLTC", "12 pr", "18")
    for _ in range(ntcx):
        add(src, dst, "Thermocouples", "TCX", "12 pr", "16")
    for _ in range(nrtd):
        add(src, dst, "RTDs", "RTD", "8 tri", "18")


# sources ---------------------------------------------------------------------------------------------------
S13 = [BY_TAG["SWGR-13.8-1"], BY_TAG["SWGR-13.8-2"], BY_TAG["SWGR-13.8-3"]]
S4 = [BY_TAG["SWGR-4.16-A"], BY_TAG["SWGR-4.16-B"]]
R1, RH, CR = BY_TAG["R1"], BY_TAG["RH"], BY_TAG["CR"]
LVS = {"A": [BY_TAG["R2A"], BY_TAG["R2B"], BY_TAG["R2C"], BY_TAG["R3"], BY_TAG["LC-480-A"], BY_TAG["LC-480-B"]],
       "B": [src_of(r"^R4 480 V switchgear")], "C": [RH], "D": [BY_TAG["MCC-C1"]], "E": [BY_TAG["MCC-C1"], BY_TAG["MCC-C2"]],
       "F": [BY_TAG["MCC-C2"]], "G": [src_of(r"^CT MCC / VFD e-house"), src_of(r"^CCS MV switchgear building")], "H": [src_of(r"^BESS auxiliary transformer")],
       "I": [src_of(r"^MOD-LV 480 V"), src_of(r"^PAD-LV"), src_of(r"^PCM: power control")], "J": [src_of(r"^Inlet-chilling e-house")],
       "K": [src_of(r"^LNG 17"), src_of(r"^H2 20")]}
CTRLS = {"A": [BY_TAG["R2A"], BY_TAG["R2B"], BY_TAG["R2C"], BY_TAG["R3"], R1], "B": [BY_TAG["R4"]], "C": [RH],
         "D": [src_of(r"^Gas yard local control")], "E": [src_of(r"Water treatment building"), R1], "F": [CR],
         "G": [src_of(r"^CCS F&G / control")], "H": [src_of(r"^BESS EMS")], "I": [src_of(r"^PCM: power control"), src_of(r"^PAD-EH")],
         "J": [src_of(r"^Inlet-chilling e-house")], "K": [src_of(r"^LNG 17"), src_of(r"^H2 20")]}


def nearest(cands, it):
    c = centre(it)
    return min(cands, key=lambda s: math.dist(centre(s), c))


def gt_no(it):
    m = re.search(r"(?:GT|GTG-|SOS-|FS-GT|WASH-GT|GFM-|TCP-|FH-|INL-|SFC-|EXC-|NGT-|SPC-|ET-|UAT-|GCB-|BFP-|HRSG[- ]|SCRB-|CEMS-)(\d)",
                  it["name"] + " " + (it.get("tag") or ""))
    return int(m.group(1)) if m else None


def lv_src(it):
    a = it["area"]
    n = gt_no(it)
    if a == "A" and n and re.search(r"GT\d|GTG|SOS|FS-GT|WASH|GFM|FH-|TCP|INL", it["name"]):
        return BY_TAG[f"MCC-GT{n}"]
    return nearest(LVS.get(a, [BY_TAG["MCC-C1"]]), it)


def ctl_src(it):
    if re.search(r"230 kV breaker|entrance|terminal tower", it["name"]):
        return RH
    return nearest(CTRLS.get(it["area"], [R1]), it)


# MV distribution and large drives ---------------------------------------------------------------------------
for n in (1, 2, 3):
    uat = BY_TAG[f"UAT-{n}"]
    mv_feeder(uat, S13[n - 1], 40000, 13.8, f"UAT-{n} secondary to SWGR-13.8-{n} incomer", motor=False,
              note="40 MVA UAT (typ.); often non-segregated bus duct instead of cable")
    mv_feeder(S13[n - 1], BY_TAG[f"TX-LCI-{n}"], 9000, 13.8, f"Static starter LCI-{n} input transformer", motor=False)
    et = BY_TAG[f"ET-{n}"]
    mv_feeder(BY_TAG[f"GSU-{n}"], et, 3000, 21, f"Excitation transformer ET-{n} from the IPB tap", motor=False,
              note="21 kV from the generator terminals (IPB tap)")
    bfp = BY_TAG[f"BFP-{n}"]
    mv_feeder(S13[n - 1], bfp, 5000, 13.8, f"Boiler feed pump BFP-{n}A / {n}B", n=2, note="2 x 100 %, ~5 MW each (typ.)")
for t, s in (("T-R2A", 0), ("T-R2B", 1), ("T-R2C", 2), ("T-R3", 1), ("LCT-A", 0), ("LCT-B", 2)):
    it = one(rf"^{t}:") or BY_TAG.get(t)
    mv_feeder(S13[s], it, 2500 if t.startswith("T-R2") else 3000, 13.8, f"{t} unit / load-centre transformer primary",
              motor=False)
for k, t in enumerate(("T4-A", "T4-B")):
    mv_feeder(S13[0 if k == 0 else 2], BY_TAG[t], 10000, 13.8, f"{t} station transformer primary", motor=False)
    mv_feeder(BY_TAG[t], S4[k], 10000, 4.16, f"{t} secondary to SWGR-4.16-{'AB'[k]} incomer", motor=False)
for k in range(4):
    tr = one(rf"^T-R4-{k + 1}:")
    mv_feeder(S13[k % 3], tr, 3750, 13.8, f"T-R4-{k + 1} ACC VFD transformer primary", motor=False)
cp = BY_TAG["CP-A/B/C"]
for k in range(3):
    mv_feeder(S4[k % 2], cp, 900, 4.16, f"Condensate pump CP-{'ABC'[k]}")
# ACC fans: 80 cells, VFD in R4 to each fan motor (motor positions spread over the ACC footprint)
acc = BY_TAG["ACC"]
r4 = BY_TAG["R4"]
fa = acc["fp"]
for k in range(80):
    i, j = k % 10, k // 10
    x = fa[0] + (fa[1] - fa[0]) * (i + .5) / 10
    y = fa[2] + (fa[3] - fa[2]) * (j + .5) / 8
    cell = dict(acc, fp=[x - 2, x + 2, y - 2, y + 2], name=f"ACC fan motor F{k + 1:02d}", tag=f"ACC-F{k + 1:02d}", id=f"accf{k}")
    route, vert, kind = route_len(r4, acc)
    extra = abs(x - (fa[0] + fa[1]) / 2) + abs(y - (fa[2] + fa[3]) / 2) + 70          # along the fan deck, up the column
    I = 150 / (1.732 * .48 * .85)
    s, sets, vd = size_lv(I, (route + vert + extra) * 1.1)
    add(r4, cell, f"ACC fan F{k + 1:02d} motor (VFD output)", "VFD", "3/C+3G", s, sets, 1, 150, 480, round(I),
        f"VD {vd:.1f} %", length=(route + extra, vert, kind))
    add(r4, cell, f"ACC fan F{k + 1:02d} vibration / RTD / local station", "CT", "7/C", "14", length=(route + extra, vert, kind))
    pa = PATHS.get((r4["id"], acc["id"]))
    if pa:
        fan = pa[:-1] + [(pa[-1][0], pa[-1][1], 70.0), (round(x, 1), round(y, 1), 70.0)]
        ROWS[-1]["path"] = ROWS[-2]["path"] = fan

# carbon capture
ccs_sw = src_of(r"^CCS MV switchgear building")
for k, t in enumerate(find(r"^CCS T-\d: ")):
    mv_feeder(t, ccs_sw, 60000, 13.8, f"{t['tag']} secondary to CCS 13.8 kV switchgear", motor=False, note="60 MVA (typ.)")
for bf in find(r"^BF-[ABC]"):
    mv_feeder(ccs_sw, bf, 16000, 13.8, f"Booster fan {bf['tag']} (MV VFD output)", armored=True, note="~16 MW per drawing")
for k in range(3):
    mv_feeder(ccs_sw, one(r"^CO2 compression"), 19000, 13.8, f"CO2 compressor train {k + 1}", note="~19 MW per drawing")
mv_feeder(ccs_sw, one(r"HP export compressor"), 8000, 13.8, "HP CO2 export compressor (conditional)")
for k in range(4):
    mv_feeder(ccs_sw, one(r"^CCS circulating-water pumps"), 1200, 4.16, f"CCS circulating-water pump {k + 1}")
for it in find(r"^Rich/lean pumps"):
    mv_feeder(ccs_sw, it, 450, 4.16, f"Rich / lean amine pumps ({it['name'][-1]})", n=2, armored=True)
mv_feeder(ccs_sw, src_of(r"^CT MCC / VFD e-house"), 3000, 13.8, "CT MCC / VFD e-house transformer", motor=False, n=2)
# fuel, utilities and options
fg = one(r"^CONDITIONAL: fuel-gas compressors")
if fg:
    mv_feeder(S13[2], fg, 3000, 13.8, "Fuel-gas compressor (conditional)", armored=True, n=3)
icx = src_of(r"^Inlet-chilling e-house")
mv_feeder(S13[0], icx, 12000, 13.8, "Inlet-chilling e-house / transformer", motor=False)
for k in range(6):
    mv_feeder(icx, one(r"^Chillers x6"), 1500, 4.16, f"Water-cooled chiller CH-{k + 1}")
for k in range(4):
    mv_feeder(icx, one(r"^CW pumps \(inlet"), 450, 4.16, f"Inlet-chilling CW pump {k + 1}")
for pat, kw, nm in ((r"^LNG 17", 5000, "LNG e-house"), (r"^H2 20", 20000, "H2 13.8 kV e-house (electrolyzers)")):
    it = one(pat)
    if it:
        mv_feeder(S13[2], it, kw, 13.8, nm, motor=False)
# BESS: DC from containers to PCS skids, 35 kV collection, MPT
pcs = find(r"^BESS PCS / MV skid")
coll = src_of(r"^34.5 kV collector e-house")
for c in find(r"^BESS container"):
    sk = nearest(pcs, c)
    add(c, sk, "DC 1,500 V: container to PCS (2 x 1/C per pole pair)", "PV", "1/C", "500", 4, 2, 4000, 1500,
        None, "4 sets of +/- per container (typ.)")
    ctrl(sk, c, n7=1)
    inst(sk, c, n4=1)
    add(sk, c, "Container auxiliary 480 V", "LV", "3/C+G", "2", 1, 1, 60, 480, 85)
pcs_sorted = sorted(pcs, key=lambda s: centre(s))
for loop in range(3):
    chain = [coll] + pcs_sorted[loop * 5:(loop + 1) * 5] + [coll]
    for a, b in zip(chain, chain[1:]):
        mv_feeder(a, b, 25000, 34.5, f"34.5 kV collection loop {loop + 1}", motor=False)
mpt = one(r"^BESS main power transformer")
mv_feeder(coll, mpt, 100000, 34.5, "BESS MPT 34.5 kV side", motor=False, note="100 MVA (typ.)")
# modular yard
mod_eh = BY_TAG["MOD-EH"]
pad_eh = BY_TAG["PAD-EH"]
cont_eh = BY_TAG["CONT-EH"]
pcm = src_of(r"^PCM: power control")
for k in range(8):
    eg = one(rf"^RICE engine-generator {k + 1}$")
    mv_feeder(eg, pcm, 18400, 13.8, f"RICE-{k + 1} generator leads to the PCM", motor=False, note="18V50 class ~18 MW (typ.)")
for k in (1, 2):
    sc = one(rf"^SC-{k}: aeroderivative")
    mv_feeder(sc, one(rf"^SC-{k} PCM"), 50000, 13.8, f"SC-{k} (LM6000 class) generator leads to its PCM / GCB", motor=False)
    mv_feeder(one(rf"^SC-{k} PCM"), mod_eh, 50000, 13.8, f"SC-{k} PCM to MOD-EH", motor=False)
BUS_DUCT = [("RICE PCM to MOD-EH bus tie", "13.8 kV, ~6,300 A"), ("MOD-EH to T-MOD-1 LV side", "13.8 kV, ~5,000 A"),
            ("MOD-EH to T-MOD-2 LV side", "13.8 kV, ~5,000 A"),
            ("Generator terminals to GSU / UAT / ET (isolated-phase bus)", "21 kV, as modelled (IPB routes)"),
            ("T-R4-1..4 secondaries to the ACC VFD lineups", "480 V, ~4,500 A each")]
mv_feeder(one(r"^MOD-LV 480 V"), mod_eh, 7500, 13.8, "MOD-LV step-up to MOD-EH", motor=False)
for it in find(r"^Fuel-gas compressors 3 x 100"):
    mv_feeder(mod_eh, it, 1500, 13.8, "Modular-yard fuel-gas compressor", armored=True, n=3)
for k in range(1, 9):
    g = BY_TAG.get(f"CONT-{k}")
    if not g:
        continue
    dst = cont_eh if k > 2 else BY_TAG.get(f"GSP-{k}")
    if k <= 2:
        add(g, dst, f"CONT-{k} 13.8 kV leads to GSP-{k}", "SHD", "3/C", "2", 1, 1, 2000, 13800, 93, "portable, on grade")
        add(dst, pad_eh, f"GSP-{k} to PAD-EH", "SHD", "3/C", "2", 1, 1, 2000, 13800, 93, "portable, on grade")
    else:
        add(g, dst, f"CONT-{k} 13.8 kV to CONT-EH (ground tray)", "MV15", "1/C", "2", 1, 3, 2000, 13800, 93)
add(cont_eh, pad_eh, "CONT-EH to PAD-EH", "MV15", "1/C", "500", 2, 3, 12000, 13800, 558)
for k in (1, 2):
    tm = BY_TAG[f"TM-{k}"]
    add(tm, pad_eh, f"TM-{k} (TM2500) generator leads to PAD-EH", "MV15", "1/C", "1000", 3, 3, 35000, 13800, 1464,
        "3 per phase for 125 % (the model shows 2 per phase: update the model)")
mob = BY_TAG["MOB-1"]
add(mob, BY_TAG["PIC"], "MOB-1 13.8 kV to PIC through the C-L01 ramp", "SHD", "3/C", "2", 1, 1, 2000, 13800, 93, "portable")
add(BY_TAG["PIC"], BY_TAG["LB"], "PIC to commissioning load bank", "SHD", "3/C", "2", 1, 1, 2000, 13800, 93, "portable")
for t, kw in (("GEN-E", 2000), ("GEN-O", 1000)):
    g = BY_TAG[t]
    I = kw / (1.732 * .48)
    sets = math.ceil(I * 1.25 / 405)
    add(g, BY_TAG["PAD-LV"], f"{t} 480 V to PAD-LV (Type W sets, cam-lock)", "W", "1/C", "4/0", sets, 4, kw, 480,
        round(I), "4 x 1/C per set incl. ground")
add(BY_TAG["PAD-LV"], BY_TAG["PAD-TX"] if "PAD-TX" in BY_TAG else one(r"^PAD-TX"), "PAD-LV to PAD-TX", "LV1C", "1/C", "500",
    12, 4, 3000, 480, 3600, "12 sets of 4 x 1/C (or bus duct)")
fc_inv = [one(r"^Fuel-cell inverter"), one(r"^FC-5..12 inverter")]
for k in range(1, 13):
    fc = one(rf"^FC-{k}: SOFC")
    lv_feeder(fc, fc_inv[0 if k <= 4 else 1], 250, f"FC-{k} SOFC output to inverter / AC cabinet", motor=False)
for it in fc_inv:
    lv_feeder(it, one(r"^MOD-LV 480 V"), 2000 if "5..12" in it["name"] else 1000, "Fuel-cell AC cabinet to MOD-LV",
              motor=False)
for k in range(1, 7):
    mt = one(rf"^MT-{k}: microturbine")
    lv_feeder(mt, one(r"^MOD-LV 480 V"), 1000, f"MT-{k} microturbine output", motor=False)
for it in find(r"^Black-start genset"):
    lv_feeder(it, src_of(r"^BS black-start switchgear"), 2000, "Black-start genset output", motor=False)


# audit additions: circuits found missing by the first audit
emcc = BY_TAG["EMCC"]
for t in ("EDG-1", "EDG-2"):
    I = 3000 / (1.732 * .48)
    add(BY_TAG[t], emcc, f"{t} 480 V output to EMCC", "LV1C", "1/C", "500", 13, 4, 3000, 480, round(I),
        "13 sets of 4 x 1/C 500 kcmil for 3 % VD (or bus duct)")
lv_feeder(BY_TAG["LC-480-A"], emcc, 400, "EMCC normal supply from LC-480-A", motor=False)
lv_feeder(LVS["A"][4], BY_TAG["CRANE"], 150, "Bridge crane conductor-bar feeder (fused disconnect, north wall)")
for t in ("CPR-1", "CPR-2"):
    lv_feeder(lv_src(BY_TAG[t]), BY_TAG[t], 10, "Cathodic-protection rectifier supply", motor=False)
    add(BY_TAG[t], BY_TAG[t], "CP anode / structure leads (allowance)", "LV", "1/C", "8", 1, 1, length=(1500, 0, "Buried"),
        note="HMWPE anode lead cable by CP vendor; allowance")
gt = one(r"^CONDITIONAL: grounding transformer")
if gt:
    mv_feeder(src_of(r"^34.5 kV collector e-house"), gt, 2000, 34.5, "Grounding transformer connection (34.5 kV)", motor=False)
the = one(r"^Turbine hall exterior electrical")
if the:
    lv_feeder(BY_TAG["LC-480-A"], the, 120, "Turbine hall roof exhausters / louvres / exterior lighting panel", motor=False)
for pat, n in ((r"^230 kV switchyard \(gravel", 40), (r"^BESS yard: surfacing", 24)):
    it = one(pat)
    if it:
        src = RH if "switchyard" in it["name"] else LVS["H"][0]
        for c in range(math.ceil(n / 10)):
            add(src, it, f"Yard lighting circuit {c + 1} ({min(10, n - c * 10)} floodlights)", "LV", "3/C+G", "8", 1, 1, 4,
                480, 6, "pole to pole, buried", length=(10 * 90, 30, "Buried"))

# 230 kV underground circuits and overhead conductors (from the routes) ---------------------------------------
dummy = lambda name, tag: dict(name=name, tag=tag, area="C", id=tag, fp=[0, 1, 0, 1], layer="ROUTES")
for k, r in enumerate([r for r in M["routes"] if r["type"] == "hv_cable" and not r["layer"].startswith("OPT_DC")]):
    L = sum(math.dist(a, b) for a, b in zip(r["points"], r["points"][1:]))
    nm = r["layer"].replace("_ROUTES", "").replace("OPT_", "")
    add(dummy(f"230 kV {nm} tie: switchyard end", f"{nm}-230-A"), dummy(f"230 kV {nm} tie: plant end", f"{nm}-230-B"),
        f"230 kV underground circuit ({nm})", "HV230", "1/C", "2500", 1, 3, None, 230000, None,
        "route length from the model + 2 x 40 ft risers", length=(L, 80, "HV duct bank"))
for k, r in enumerate([r for r in M["routes"] if r["type"] in ("hv_overhead", "hmod")]):
    L = sum(math.dist(a, b) for a, b in zip(r["points"], r["points"][1:]))
    add(dummy("230 kV overhead span start", f"OH-{k + 1}-A"), dummy("230 kV overhead span end", f"OH-{k + 1}-B"),
        f"230 kV overhead {'H-MOD tie' if r['type'] == 'hmod' else 'GSU / line'} conductors", "ACSR", "1/C", "1590", 1, 3,
        None, 230000, None, "plus 3 % sag", length=(L * 1.03, 0, "Overhead"))


# Device-level wiring: each item's field devices (typical counts for the equipment class): motors, motor-operated
# valves, analog instruments, discrete devices, solenoids, thermocouples, RTDs. Every device gets its own cable to a
# local junction box (or its MCC for motors / MOVs); junction boxes home-run to the control building on multipair.
# m = [(motors, kW each, VFD)], mov = MOVs, ai / di = analog / discrete instruments, sol, tc, rtd.
DEV = [
    (r"^GT\d: H-class", dict(ai=120, di=80, sol=30, tc=60)),
    (r"^GTG-\d", dict(m=[(4, 2, 0)], ai=20, di=20, rtd=24)),
    (r"^GT\d aux", dict(m=[(12, 45, 0)], ai=40, di=40, sol=10)),
    (r"^SOS-", dict(m=[(3, 15, 0)], ai=15, di=15, sol=4)),
    (r"^FS-GT", dict(ai=4, di=20, sol=8)),
    (r"^WASH-GT", dict(m=[(2, 22, 0)], ai=4, di=6, sol=4)),
    (r"^GFM-", dict(ai=15, di=15, sol=6)),
    (r"^FH-\d:", dict(m=[(2, 15, 0)], ai=10, di=30, sol=60)),
    (r"^INL-|inlet duct", dict(ai=6, tc=4)),
    (r"^SFC-|^EXC-|^NGT-|^SPC-", dict(ai=10, di=20)),
    (r"^GSU-|^UAT-|^T4-|^LCT-|^T-R\d|^T-R4|^CCS T-\d:|main power transformer|^T-MOD", dict(m=[(6, 2, 0)], ai=8, di=16)),
    (r"^GCB-", dict(di=20)),
    (r"^HRSG \d \+", dict(m=[(4, 15, 0)], mov=45, ai=160, di=80, sol=20, tc=80)),
    (r"^BFP-", dict(m=[(4, 15, 0)], ai=30, di=20, rtd=24)),
    (r"^SCRB-", dict(m=[(2, 110, 0)], ai=6, di=6)),
    (r"^CEMS shelter", dict(m=[(2, 5, 0)], ai=12, di=10)),
    (r"blowdown tank", dict(mov=4, ai=6, di=4)),
    (r"chemical feed|^Chemical feed", dict(m=[(6, 2, 0)], ai=8, di=12)),
    (r"^SWAS", dict(ai=24, di=8)),
    (r"^ST: steam", dict(mov=20, ai=150, di=100, sol=20, tc=40)),
    (r"^STG", dict(m=[(4, 2, 0)], ai=20, di=20, rtd=30)),
    (r"^ST aux", dict(m=[(10, 40, 0)], ai=30, di=30)),
    (r"^EHC-ST", dict(m=[(2, 37, 0)], ai=10, di=10, sol=10)),
    (r"^GSC-ST", dict(m=[(2, 30, 0)], ai=6, di=6)),
    (r"^Air-cooled condenser", dict(mov=30, ai=100, di=160)),
    (r"^Aux dry coolers", dict(m=[(12, 22, 1)], ai=12, di=24)),
    (r"^CCW-A/B", dict(m=[(2, 150, 0)], ai=10, di=10)),
    (r"^CP-A/B/C", dict(ai=15, di=15, rtd=18)),
    (r"^VAC-A/B", dict(m=[(2, 75, 0)], ai=10, di=10)),
    (r"^Air compressors", dict(m=[(3, 250, 0)], ai=15, di=15)),
    (r"tank|Tank", dict(ai=4, di=4)),
    (r"230 kV breaker", dict(ai=2, di=20)),
    (r"entrance: surge|terminal tower", dict(di=4)),
    (r"filter-separator", dict(ai=8, di=12, sol=4)),
    (r"Gas metering|M&R 5|gas quality", dict(ai=20, di=8)),
    (r"regulation", dict(ai=12, di=12, sol=6)),
    (r"performance heater", dict(ai=10, di=10)),
    (r"ESD valve|insulating joint", dict(di=8, sol=4)),
    (r"Auxiliary boiler", dict(m=[(3, 30, 0)], ai=25, di=25)),
    (r"Ammonia storage", dict(m=[(2, 7, 0)], ai=8, di=10)),
    (r"Fire pump house", dict(m=[(1, 250, 0), (1, 15, 0)], ai=10, di=30)),
    (r"^Wastewater treatment", dict(m=[(8, 15, 0)], ai=20, di=30)),
    (r"Oil-water", dict(m=[(2, 7, 0)], ai=4, di=4)),
    (r"Water treatment building", dict(m=[(20, 15, 0)], ai=60, di=80)),
    (r"H2 / CO2 storage", dict(ai=8, di=8)),
    (r"Stormwater pump", dict(m=[(2, 75, 0)], ai=4, di=6)),
    (r"LS-1", dict(m=[(2, 7, 0)], ai=2, di=4)),
    (r"^Absorber [ABC] \(", dict(mov=10, ai=60, di=40)),
    (r"intercooler", dict(m=[(2, 150, 0)], ai=15, di=15)),
    (r"^DCC-[ABC]:", dict(ai=20, di=10)),
    (r"^DCC-[ABC] pumps", dict(m=[(2, 200, 0)], ai=10, di=10)),
    (r"^Water-wash pumps", dict(m=[(2, 110, 0)], ai=8, di=8)),
    (r"^STR-", dict(mov=6, ai=50, di=30)),
    (r"^RB-", dict(mov=4, ai=15, di=10)),
    (r"^Rich/lean", dict(ai=20, di=20, rtd=12)),
    (r"^DMP-", dict(m=[(2, 15, 0)], ai=8, di=16)),
    (r"^BF-[ABC]", dict(m=[(4, 15, 0)], ai=30, di=20, rtd=18)),
    (r"^CO2 compression", dict(m=[(10, 30, 0)], ai=150, di=100, rtd=36)),
    (r"^CO2 dehydration", dict(m=[(2, 20, 0), (1, 200, 0)], mov=12, ai=30, di=30)),
    (r"^CCS-CT|CCS cooling tower", dict(m=[(30, 110, 1)], ai=30, di=60)),
    (r"^CCS circulating", dict(ai=20, di=20)),
    (r"instrument air", dict(m=[(2, 250, 0)], ai=10, di=10)),
    (r"^Reclaimer", dict(m=[(2, 15, 0)], ai=15, di=10)),
    (r"Solvent", dict(m=[(4, 15, 0)], ai=10, di=10)),
    (r"^CCS wastewater", dict(m=[(6, 15, 0)], ai=15, di=15)),
    (r"Activated-carbon", dict(ai=6, di=6)),
    (r"HP export compressor", dict(m=[(4, 30, 0)], ai=60, di=40, rtd=12)),
    (r"^BESS container", dict(di=8)),
    (r"^BESS PCS", dict(m=[(2, 7, 0)], ai=6, di=10)),
    (r"^RICE engine-generator", dict(m=[(6, 22, 0)], ai=80, di=60, tc=30, rtd=12)),
    (r"^RICE \d radiators", dict(m=[(5, 30, 1)], ai=5, di=10)),
    (r"^RICE \d SCR", dict(m=[(1, 15, 0)], ai=10, di=6)),
    (r"^SC-\d: aeroderivative", dict(m=[(10, 30, 0)], ai=150, di=100, tc=40)),
    (r"^SC-\d lube-oil fin-fan", dict(m=[(4, 15, 0)], ai=4, di=8)),
    (r"^SC-\d SCR", dict(m=[(2, 30, 0)], ai=12, di=8)),
    (r"^TM-\d", dict(m=[(6, 22, 0)], ai=100, di=60, tc=20)),
    (r"^CONT-\d|^MOB-1|^GEN-[EO]", dict(ai=20, di=20)),
    (r"^Black-start genset", dict(ai=20, di=20)),
    (r"^FC-\d+:", dict(ai=10, di=10)),
    (r"^MT-\d", dict(ai=6, di=6)),
    (r"^Fuel-gas compressors 3|CONDITIONAL: fuel-gas compressors", dict(m=[(6, 15, 0)], ai=30, di=20, rtd=12)),
    (r"^Chillers x6", dict(m=[(6, 7, 0)], ai=120, di=120)),
    (r"^CW pumps \(inlet", dict(ai=24, di=24)),
    (r"^Inlet-chilling tower", dict(m=[(12, 75, 1)], ai=12, di=24)),
    (r"^CHW pumps", dict(m=[(4, 150, 0)], ai=12, di=12)),
    (r"^LNG \d+:|^H2 \d+:", dict(m=[(1, 15, 0)], ai=12, di=12, sol=2)),
]
DEVRX = [(re.compile(p), d) for p, d in DEV]
JB_PAIRS = 24
FIELD = dict(ai=("PLTC", "1 pr", "16", "Analog instrument to JB"), di=("PLTC", "1 pr", "16", "Discrete device to JB"),
             sol=("CT", "2/C", "14", "Solenoid to JB"), tc=("TCX", "1 pr", "16", "Thermocouple to TC JB"),
             rtd=("RTD", "1 tri", "18", "RTD to JB"))
N_DEV = defaultdict(int)


def dev_of(it):
    for rx, d in DEVRX:
        if rx.search(it["name"]):
            return d
    return None


def device_wiring(it, d):
    f = it["fp"]
    span = (f[1] - f[0]) + (f[3] - f[2])
    h = it["z"][1]
    field = min(250, max(40, span * .5 + h * .4 + 25))           # average device-to-JB run on the equipment
    haz = bool(HAZ.search(it["name"]))
    src_lv = lv_src(it)
    src_ct = ctl_src(it)
    tag = it.get("tag") or it["name"][:18]
    for (n, kw, vfd) in d.get("m", []):
        for k in range(n):
            N_DEV["motors"] += 1
            if vfd:
                route, vert, kind = route_len(src_lv, it)
                I = kw / (1.732 * .48 * .85)
                s, sets, vd = size_lv(I, (route + vert + field) * 1.1)
                add(src_lv, it, f"Motor {k + 1} of {n}, {kw} kW (VFD output)", "VFD", "3/C+3G", s, sets, 1, kw, 480, round(I),
                    f"VD {vd:.1f} %", length=(route + field, vert, kind))
            else:
                lv_feeder(src_lv, it, kw, f"Motor {k + 1} of {n}, {kw} kW")
            add(src_lv, it, f"Motor {k + 1} of {n}: control / local station", "CTAX" if haz else "CT", "7/C", "14")
    for k in range(d.get("mov", 0)):
        N_DEV["MOVs"] += 1
        lv_feeder(src_lv, it, 2.5, f"MOV {k + 1}: power", prod="LVAX" if haz else "LV")
        add(src_lv, it, f"MOV {k + 1}: control", "CTAX" if haz else "CT", "12/C", "14")
    for key in ("ai", "di", "sol", "tc", "rtd"):
        n = d.get(key, 0)
        if not n:
            continue
        N_DEV[key] += n
        prod, cond, size, what = FIELD[key]
        if key in ("ai", "di") and haz:
            prod = "ITC"
        if key == "sol" and haz:
            prod = "CTAX"
        per = 10 if key in ("sol", "rtd") else 20
        for j in range(math.ceil(n / per)):
            m_ = min(per, n - j * per)
            jb = f"JB-{tag}-{key.upper()}{j + 1:02d}"
            add(it, dict(it, tag=jb), f"{what} {jb}: {m_} x {cond} (field cables)", prod, cond, size, m_, 1,
                note="device-to-JB field wiring, one cable per device", length=(field, 10, "Tray / conduit"))
            hp = {"ai": ("ITC" if haz else "PLTC", "24 pr", "18"), "di": ("ITC" if haz else "PLTC", "24 pr", "18"),
                  "tc": ("TCX", "24 pr", "16"), "rtd": ("RTD", "12 tri", "18"),
                  "sol": ("CTAX" if haz else "CT", "25/C", "14")}[key]
            add(src_ct, dict(it, tag=jb), f"Home run {jb} to the control / marshalling room", hp[0], hp[1], hp[2])

# LV power, control, instrumentation, FA, fibre per item, from its wiring applications --------------------------
DEF_KW = [(r"GT\d aux", 400), (r"GT1:|GT2:|GT3:", 150), (r"HRSG \d \+", 150), (r"FH-\d:", 40), (r"SOS-", 30),
          (r"FS-GT", 10), (r"WASH-GT", 30), (r"GFM-", 15), (r"TCP-", 15), (r"SCRB-", 2 * 110), (r"AC-A/B/C|Air compressor", 3 * 250),
          (r"CCW-A/B", 2 * 150), (r"AUXC", 12 * 22), (r"VAC-A/B", 2 * 75), (r"ST aux", 250), (r"EHC", 45), (r"GSC-ST", 45),
          (r"chemical feed|Chemical feed", 15), (r"blowdown tank", 10), (r"CEMS", 30), (r"SWAS", 15), (r"EXC-", 60),
          (r"CRANE|Bridge crane", 150), (r"intercooler", 2 * 150), (r"Water-wash pumps", 2 * 110), (r"DCC-[ABC] pumps", 2 * 200),
          (r"CCS-CT|cooling tower", 30 * 110), (r"DMP-", 20), (r"RB-", 10), (r"STR-", 30), (r"ABS-|Absorber [ABC] \(", 40),
          (r"instrument air", 2 * 250), (r"dehydration", 150), (r"Reclaimer", 75), (r"Activated-carbon", 15),
          (r"Solvent|NaOH", 20), (r"wastewater|Wastewater", 60), (r"Fire pump house", 60), (r"Auxiliary boiler", 75),
          (r"Ammonia storage", 15), (r"Oil-water", 15), (r"tank", 10), (r"HTP-", 120), (r"LS-1|lift station", 30),
          (r"Stormwater pump", 2 * 75), (r"ULSD unloading and forwarding pumps", 2 * 55), (r"ULSD truck", 15),
          (r"performance heater", 60), (r"Gas metering|regulation|filter-separators|ESD valve", 10),
          (r"local control enclosure", 15), (r"CPR-", 10), (r"Gas conditioning", 30), (r"Aqueous ammonia", 10),
          (r"inlet chiller", 500), (r"LB:", 0), (r"Comms tower", 10), (r"Main gate", 5), (r"fence.*CCTV|CCTV", 20),
          (r"Site lighting", 60), (r"Gas yard", 10)]
BLD = re.compile(r"building|house|Gatehouse|Warehouse|workshop|e-house|R1:|R2[ABC]:|R3:|R4:|RH:|shelter|control room|"
                 r"control / admin|PCM|MOD-EH|PAD-EH|CONT paralleling|collector e-house|EMS|F&G", re.I)
SKIP_LV = re.compile(r"^R1 |SWGR|LC-480|MCC|EMCC|transformer|T-R\d|T-R4|LCT-|T4-|UAT|GSU|ET-\d|TX-LCI|VFD lineup|"
                     r"BFP|CP-A/B/C|BF-[ABC]|CO2 compression|export compressor|circulating-water|Rich/lean|CONDITIONAL: fuel-gas|"
                     r"Chillers x6|CW pumps \(inlet|BESS container|BESS PCS|BESS blocks|RICE engine-generator|SC-\d:|CONT-\d|"
                     r"TM-\d|MOB-1|GEN-[EO]|FC-\d|MT-\d|Black-start genset|Fuel-cell inverter|FC-5\.\.12|^PAD-|^PIC|^GSP|"
                     r"Fuel-gas compressors 3|Air-cooled condenser|GCB-|NGT-|SPC-|breaker|entrance|terminal tower|"
                     r"cable termination structure|GTG-|STG|230 kV switchyard|BESS yard|INL-|stair tower|Turbine hall exterior")
done_bld = set()
for it in ITEMS:
    apps = " ".join(it.get("wiring", []))
    d = dev_of(it)
    if d and it["area"] != "L":
        device_wiring(it, d)
    if not apps:
        continue
    nm = it["name"]
    a = it["area"]
    if a in ("L",):
        continue
    is_bld = bool(BLD.search(nm)) and not re.search(r"skid|pump|^R1 ", nm)
    # LV power
    if not SKIP_LV.search(nm) and re.search(r"LV power|LV 480", apps) and not (d and d.get("m")):
        kw = next((k for p, k in DEF_KW if re.search(p, nm)), None)
        if kw is None and is_bld:
            f = it["fp"]
            kw = max(30, round((f[1] - f[0]) * (f[3] - f[2]) * .012))          # ~12 W / ft2 lighting, HVAC, small power
        if kw is None:
            kw = 15
        if kw > 0:
            lv_feeder(lv_src(it), it, kw, "LV power feeder" if not is_bld else "LV panel / distribution feeder")
    # control and instrumentation
    big = re.search(r"GT\d:|STG|GTG|HRSG \d \+|ST: steam|ACC|CO2 compression|BF-|SC-\d:|RICE engine|TM-\d", nm)
    if ("Control" in apps or "CT" in apps) and (not d or re.search(r"230 kV breaker|GCB-|GSU|UAT|transformer", nm)):
        if re.search(r"230 kV breaker", nm):
            ctrl(RH, it, n7=0, n12=4, nct=3)
        elif re.search(r"GCB-|GSP-", nm):
            ctrl(ctl_src(it), it, n7=0, n12=3, nct=4)
        elif re.search(r"GSU|UAT|T-|transformer|T4-|LCT", nm):
            ctrl(ctl_src(it), it, n7=1, n12=1, nct=2)
        elif big:
            ctrl(ctl_src(it), it, n7=2, n12=6)
        elif not re.search(r"^R1 |SWGR|LC-480|MCC", nm):
            ctrl(ctl_src(it), it, n7=1, n12=0)
    if "Instrumentation" in apps and not d:
        if big:
            inst(ctl_src(it), it, n4=2, n12=8)
        elif re.search(r"skid|pumps|compressor|tank|heater|separator|absorber|stripper|DCC|reboiler|module", nm, re.I):
            inst(ctl_src(it), it, n4=1, n12=1)
        else:
            inst(ctl_src(it), it, n4=1)
    if "Thermocouple" in apps and not d:
        inst(ctl_src(it), it, n4=0, ntcx=6 if big else 2)
    if "RTD" in apps and not d:
        inst(ctl_src(it), it, n4=0, nrtd=4)
    if "Vibration" in apps:
        add(ctl_src(it), it, "Vibration monitoring", "PLTC", "4 pr", "18")
    if "Fire alarm" in apps or "Fire and gas" in apps:
        add(ctl_src(it), it, "Fire alarm / F&G loop (to the local FACP / F&G node)", "FA", "2/C", "14")
    if "Data" in apps and (is_bld or big):
        add(CR, it, "Plant network / DCS fibre", "FO", "12 SM", "-")


# outdoor and structure lighting, 480 V welding receptacles, heat-trace circuits (allowances counted from the model)
LIGHT_STRUCT = [(r"^HRSG \d \+", 40), (r"^Air-cooled condenser", 80), (r"^Absorber [ABC] \(", 30), (r"^STR-", 20),
                (r"^FH-\d:", 12), (r"stair tower", 8), (r"^DCC-[ABC]:", 8), (r"CCS cooling tower", 30),
                (r"Inlet-chilling tower", 12), (r"^HRSG \d stack|HRSG \d stack", 6), (r"RICE \d SCR", 6), (r"^SC-\d stack", 4)]
for it in ITEMS:
    for p_, n in LIGHT_STRUCT:
        if re.search(p_, it["name"]):
            src = lv_src(it)
            for c in range(math.ceil(n / 12)):
                rl, rv, rk = route_len(src, it)
                sz, st_, vd = size_lv(5, (rl + rv + 300) * 1.1)
                add(src, it, f"Area / platform lighting circuit {c + 1} ({min(12, n - c * 12)} fixtures)", "LV", "3/C+G",
                    max(sz, "10", key=SIZES_LV.index), st_, 1, 3, 480, 5, "fixture-to-fixture run included",
                    length=(rl + 300, rv, rk))
            break
site_l = next((it for it in ITEMS if it["name"].startswith("Site lighting")), None)
if site_l:
    poles = sum(1 for p_ in M["parts"] if p_["item"] == site_l["id"] and p_["kind"] == "rod" and abs(p_["a"][2] - p_["b"][2]) > 20)
    poles = poles or 120
    add(site_l, site_l, f"Site / road lighting: {poles} poles, buried circuits pole to pole", "LV", "3/C+G", "8", poles, 1,
        None, 480, None, "~120 ft between poles + riser", length=(120, 30, "Buried"))
welds = 0
for it in ITEMS:
    if re.search(r"^HRSG \d \+|^GT\d: H-class|^ST: steam|^Air-cooled condenser|^Absorber [ABC] \(|^STR-|"
                 r"^CO2 compression|^RICE engine-generator|^BF-", it["name"]):
        for k in range(4):
            welds += 1
            lv_feeder(lv_src(it), it, 60, f"480 V welding receptacle {k + 1}", motor=False)
for it in ITEMS:
    if it["name"].startswith("HTP-"):
        for k in range(24):
            add(it, it, f"Heat-trace circuit {k + 1} (power to the trace junction box)", "LV", "3/C+G", "10", 1, 1, 4, 480, 6,
                "heat-trace cable itself by others", length=(180, 20, "Tray"))

# lighting / small power branch circuits (allowance per building by floor area)
for it in ITEMS:
    if "Lighting" in " ".join(it.get("wiring", [])) and BLD.search(it["name"]) and it["name"] not in done_bld:
        done_bld.add(it["name"])
        f = it["fp"]
        area = (f[1] - f[0]) * (f[3] - f[2])
        add(it, it, "Lighting, receptacles and small power branch circuits (allowance)", "THHN", "1/C", "12", 1, 1,
            None, 277, None, "allowance: 1.6 ft of conductor per ft2 of floor", length=(area * 1.6, 0, "Conduit"))
# grounding: the modelled ground grid + equipment ground leads
gg = next(it for it in M["items"] if "station ground grid" in it["name"])
L = sum(math.dist(p["a"], p["b"]) for p in M["parts"] if p["item"] == gg["id"] and p["kind"] == "rod")
add(dict(gg, area="C"), dict(gg, area="C"), "Station ground grid, rods and risers (from the model)", "BARE", "1/C", "4/0", 1, 1,
    None, None, None, "measured from the modelled grid", length=(L, 0, "Buried"))
n_eq = sum(1 for it in ITEMS if "Grounding" in " ".join(it.get("wiring", []))) + N_DEV["motors"]
add(dict(gg, area="C", name="Equipment grounding"), dict(gg, area="C", name="Equipment grounding"),
    f"Equipment ground leads: {n_eq} items x 2 leads x 30 ft", "BARE", "1/C", "4/0", 1, 1, None, None, None, "allowance",
    length=(n_eq * 60, 0, "Buried"))



def audit():
    """Independent checks on the schedule; returns [(check, result, count, examples)]."""
    out = []
    by_dst = defaultdict(list)
    for r in ROWS:
        by_dst[r["dst"]["id"]].append(r)
    srcs = {r["src"]["id"] for r in ROWS}
    power = ("LV", "LVAX", "VFD", "LV1C", "MV15", "MV15AX", "MV5", "MV5S", "MV25", "MV35", "W", "SHD", "PV")
    miss = [it["name"][:50] for it in ITEMS if re.search(r"LV power|MV power|LV 480", " ".join(it.get("wiring", [])))
            and not any(r["prod"] in power for r in by_dst.get(it["id"], [])) and it["id"] not in srcs
            and not re.search(r"VFD lineup|BESS blocks", it["name"])]
    out.append(("Every powered item has a power circuit (VFD lineups: bus duct; BESS blocks: per container)",
                "PASS" if not miss else "CHECK", len(miss), "; ".join(miss[:4])))
    miss = [it["name"][:50] for it in ITEMS if "Instrumentation" in " ".join(it.get("wiring", []))
            and not any(r["prod"] in ("PLTC", "ITC", "TCX", "RTD") for r in by_dst.get(it["id"], []))]
    out.append(("Every instrumented item has instrument cables", "PASS" if not miss else "CHECK", len(miss), "; ".join(miss[:4])))
    dbl = [rows[0]["dst"]["name"][:40] for rows in by_dst.values()
           if any(r["service"].startswith("LV power feeder") for r in rows) and any(r["service"].startswith("Motor ") for r in rows)]
    out.append(("No lump feeder where motors are fed individually (double count)", "PASS" if not dbl else "FAIL", len(dbl),
                "; ".join(dbl[:4])))
    bad = []
    for r in ROWS:
        I = r["amps"]
        if not I:
            continue
        if r["prod"] in ("LV", "LVAX", "VFD", "LV1C") and r["size"] in AMP_LV and AMP_LV[r["size"]] * r["sets"] < 1.25 * I * .999:
            bad.append(r["service"][:40])
        if r["prod"] in ("MV15", "MV15AX", "MV5", "MV5S", "MV25") and AMP_MV[r["size"]] * .85 * r["sets"] < 1.25 * I * .999:
            bad.append(r["service"][:40])
        if r["prod"] == "MV35" and AMP_MV35[r["size"]] * .85 * r["sets"] < 1.25 * I * .999:
            bad.append(r["service"][:40])
    out.append(("Ampacity >= 125 % of full-load current (recomputed)", "PASS" if not bad else "FAIL", len(bad), "; ".join(bad[:4])))
    vdb = []
    for r in ROWS:
        if r["prod"] in ("LV", "LVAX", "VFD", "LV1C") and r["amps"] and r["size"] in R_LV:
            L = (r["route"] + r["vert"]) * 1.1
            vd = 1.732 * r["amps"] * L / 1000 * R_LV[r["size"]] / r["sets"] / 480 * 100
            if vd > 3.05:
                vdb.append(f"{r['service'][:30]} {vd:.1f}%")
    out.append(("480 V voltage drop <= 3 % (recomputed)", "PASS" if not vdb else "CHECK", len(vdb), "; ".join(vdb[:4])))
    haz = [r["dst"]["name"][:40] for r in ROWS if HAZ.search(r["dst"]["name"]) and r["prod"] in ("LV", "CT", "PLTC")
           and not re.match(r"GT\d:", r["dst"]["name"])]
    out.append(("Class I Div 2 loads on ARMOR-X / IS cable", "PASS" if not haz else "FAIL", len(haz), "; ".join(haz[:4])))
    est = [r for r in ROWS if r["kind"] == "Estimated"]
    out.append(("Runs that follow a modelled route", "INFO", len(ROWS) - len(est),
                f"{len(est)} estimated (build-out bays D4-D6 and utilities without drawn trench / bank)"))
    lng = [r for r in ROWS if r["route"] + r["vert"] > 3000 and r["prod"] not in ("FO", "THHN", "BARE", "ACSR")]
    out.append(("Copper runs > 3,000 ft", "INFO", len(lng), "; ".join(f"{r['service'][:30]}" for r in lng[:4])))
    many = [r for r in ROWS if r["sets"] >= 6 and r["prod"] in ("LV", "LVAX", "LV1C", "MV15", "MV5S", "MV25", "MV35")]
    out.append(("Feeders with >= 6 parallel sets (bus-duct candidates)", "INFO", len(many),
                "; ".join(f"{r['service'][:30]} x{r['sets']}" for r in many[:4])))
    out.append(("Excluded as bus duct", "INFO", len(BUS_DUCT), "; ".join(a for a, b in BUS_DUCT)))
    return out


XCLS = ["HV 230 kV", "MV 35 kV", "MV 25 kV", "MV 15 kV", "MV 5 kV", "LV 600 V", "DC 2 kV", "Portable MV", "Portable LV",
        "Control", "Instrumentation", "Fire alarm", "Fibre / data"]


def export_xray():
    """Compact cable data for the viewer's cable X-ray: unique paths (ints, ft) + one record per schedule row."""
    paths, pidx, recs = [], {}, []
    seg = defaultdict(int)
    for r in ROWS:
        cl = CLASS[r["prod"]]
        if cl not in XCLS:
            continue
        pi = -1
        if r.get("path") and len(r["path"]) >= 2:
            key = tuple(r["path"])
            if key not in pidx:
                pidx[key] = len(paths)
                paths.append([int(round(v)) for q in r["path"] for v in q])
            pi = pidx[key]
            n = r["sets"] * r["cps"]
            for a_, b_ in zip(r["path"], r["path"][1:]):
                k = (tuple(int(round(v)) for v in a_), tuple(int(round(v)) for v in b_))
                seg[tuple(sorted(k))] += n
        run = math.ceil((r["route"] + r["vert"] + 20) * 1.1 / 10) * 10
        recs.append([r["no"], XCLS.index(cl), r["prod"], AWG(r["size"]) if r["size"] != "-" else "", r["cond"],
                     r["sets"] * r["cps"], run, (r["src"].get("tag") or r["src"]["name"][:40]),
                     (r["dst"].get("tag") or r["dst"]["name"][:40]), r["src"]["id"], r["dst"]["id"],
                     (r["service"] if not r["desc"] else r["desc"])[:80], pi])
    fill = [[*a, *b, n] for (a, b), n in seg.items() if a != b]
    data = dict(classes=XCLS, products={k: P[k] for k in P}, paths=paths, cables=recs, fill=fill)
    json.dump(data, open(os.path.join(HERE, "brochure", "cables_xray.json"), "w"), separators=(",", ":"))
    return len(recs), len(paths)

# ------------------------------------------------------------------------------------------------ workbook
def write():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    F = "Arial"
    hdr_fill = PatternFill("solid", fgColor="1F2A33")
    hdr_font = Font(name=F, bold=True, color="FFFFFF", size=9)
    inp = Font(name=F, color="0000FF", size=9)
    blk = Font(name=F, size=9)
    grn = Font(name=F, color="008000", size=9)
    yel = PatternFill("solid", fgColor="FFFF00")
    thin = Border(bottom=Side(style="thin", color="D0D0D0"))

    # Read Me / assumptions
    ws = wb.active
    ws.title = "Read Me"
    lines = [
        ("SK-3X1 Rev 14 gas power campus: cable schedule and cable bill of materials", True),
        ("Conceptual, derived from the SK-3X1 3D model. Not engineered, not for construction or procurement.", False),
        ("", False),
        ("How it was built", True),
        ("1. Every powered item of the model (BTM data-centre option excluded) is matched to an equipment class that gives "
         "its power circuits (voltage, typical rating, source bus) and its field devices (motors, MOVs, analog and discrete "
         "instruments, solenoids, thermocouples, RTDs; typical counts for the class). Each device has its own cable to its "
         "MCC or to a local junction box; junction boxes home-run on 24-pair / 12-triad / 25-conductor cables. Device-to-JB "
         "cables are grouped per JB (Sets = number of device cables). Ratings are typical unless the drawing states one.",
         False),
        ("2. Lengths follow the model's electrical routes (trays, duct banks, buried / surface cable, trenches) as a graph: "
         "shortest path source to load, plus drops and vertical transitions. 230 kV circuits, overhead conductors and the "
         "ground grid are measured directly from the model.", False),
        ("3. Sizes: ampacity at 125 % of full-load current (NEC 310.16 75 C Cu for 600 V; typical ICEA values for MV-105), "
         "3 % voltage drop at 480 V, parallel sets above 500 kcmil (LV) / 1000 kcmil (MV).", False),
        ("4. Products are typical Southwire families for each application; confirm constructions, sizes and listings "
         "against current Southwire specifications.", False),
        ("", False),
        (f"5. {sum(1 for r in ROWS if r['kind'] == 'Estimated')} of {len(ROWS)} runs could not follow a modelled route "
         "(Route type 'Estimated'): L-shaped length between the items plus 40 ft.", False),
        ("Inputs (blue, yellow cells; every length and quantity recalculates)", True),
    ]
    for r, (t, b) in enumerate(lines, 1):
        ws.cell(r, 1, t).font = Font(name=F, bold=b, size=12 if r == 1 else 10)
    ws["A12"], ws["B12"] = "Routing slack (training, sag, cutting)", .10
    ws["A13"], ws["B13"] = "Termination allowance per run end (ft)", 10
    ws["A14"], ws["B14"] = "BOM waste / cut allowance", .05
    for c in ("B12", "B13", "B14"):
        ws[c].font = Font(name=F, color="0000FF", bold=True)
        ws[c].fill = yel
    ws["B12"].number_format = ws["B14"].number_format = "0%"
    ws["C12"] = "Source: typical industry allowance (assumption)"
    ws["C13"] = "Source: typical (assumption)"
    ws["C14"] = "Source: typical (assumption)"
    ws["A23"] = "Excluded from the cable BOM (bus duct, by others)"
    ws["A23"].font = Font(name=F, bold=True)
    for k_, (a_, b_) in enumerate(BUS_DUCT, 24):
        ws.cell(k_, 1, f"{a_}: {b_}").font = Font(name=F, size=10)
    ws["A16"] = "Sheets"
    ws["A16"].font = Font(name=F, bold=True)
    for r, t in enumerate(["Cable Schedule: one row per cable run (from / to, service, product, size, sets, length)",
                           "Cable BOM: totals per product and size, waste, order length, reels (SUMIFS on the schedule)",
                           "Accessories: terminations, glands, lugs, splices (from the schedule)",
                           "Summary: cable length by zone and voltage class",
                           "Load List: the equipment, assumed ratings and sources behind the power circuits"], 17):
        ws.cell(r, 1, t).font = Font(name=F, size=10)
    ws.column_dimensions["A"].width = 120
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 45
    for r in range(1, 25):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")

    # Cable Schedule
    sc = wb.create_sheet("Cable Schedule")
    cols = ["Cable No.", "Zone", "Area", "Service / circuit", "From (tag)", "From (description)", "To (tag)",
            "To (description)", "Class", "Product", "Product key", "Conductors", "Size (AWG/kcmil)", "Sets", "Cables / set",
            "Voltage (V)", "Load (kW)", "Current (A)", "Route type", "Route (ft)", "Drops + vertical (ft)",
            "Run length (ft)", "Total cable (ft)", "MV term. kits", "Glands / connectors", "Notes", "Cables (count)"]
    for c, h in enumerate(cols, 1):
        x = sc.cell(1, c, h)
        x.font, x.fill = hdr_font, hdr_fill
        x.alignment = Alignment(wrap_text=True, vertical="center")
    ROWS.sort(key=lambda r: (r["dst"]["area"] or "Z", CLASS[r["prod"]], r["dst"]["name"]))
    cnt = defaultdict(int)
    code = dict(MV15="MV", MV15AX="MV", MV5="MV", MV5S="MV", MV25="MV", MV35="MV", HV230="HV", LV="P", LVAX="P", VFD="P", LV1C="P",
                CT="C", CTAX="C", CTCT="C", PLTC="I", ITC="I", TCX="I", RTD="I", FA="FA", FO="F", PV="DC", SHD="PP", W="PP",
                THHN="L", BARE="G", ACSR="OH")
    for i, r in enumerate(ROWS, 2):
        a = r["dst"]["area"] or "C"
        cnt[(a, code[r["prod"]])] += 1
        no = f"{a}-{code[r['prod']]}-{cnt[(a, code[r['prod']])]:04d}"
        r["no"] = no
        key = f"{r['prod']}|{r['cond']}|{r['size']}"
        vals = [no, a, AREA.get(a, ""), r["service"] if not r["desc"] else r["desc"], r["src"].get("tag") or "",
                r["src"]["name"][:70], r["dst"].get("tag") or "", r["dst"]["name"][:70], CLASS[r["prod"]],
                P[r["prod"]], key, r["cond"], AWG(r["size"]) if r["size"] != "-" else "-", r["sets"], r["cps"],
                r["volt"], r["kw"], r["amps"], r["kind"], r["route"], r["vert"]]
        for c, v in enumerate(vals, 1):
            x = sc.cell(i, c, v)
            x.font = blk
        sc.cell(i, 22, f"=ROUNDUP((T{i}+U{i}+2*'Read Me'!$B$13)*(1+'Read Me'!$B$12),-1)").font = blk
        sc.cell(i, 23, f"=V{i}*N{i}*O{i}").font = blk
        sc.cell(i, 24, f'=IF(LEFT(I{i},2)="MV",2*N{i}*IF(L{i}="1/C",O{i},1),0)').font = blk
        sc.cell(i, 25, f'=IF(OR(LEFT(K{i},4)="LVAX",LEFT(K{i},6)="MV15AX",LEFT(K{i},4)="CTAX",LEFT(K{i},2)="W|"),'
                       f'2*N{i}*O{i},0)').font = blk
        sc.cell(i, 26, r["note"]).font = blk
        sc.cell(i, 27, f"=N{i}*O{i}").font = blk
        for c in (20, 21, 22, 23):
            sc.cell(i, c).number_format = "#,##0"
    last = len(ROWS) + 1
    widths = [12, 6, 20, 46, 12, 34, 12, 34, 14, 52, 22, 10, 12, 6, 7, 9, 9, 9, 13, 9, 10, 10, 12, 8, 9, 30, 9]
    for c, w in enumerate(widths, 1):
        sc.column_dimensions[get_column_letter(c)].width = w
    sc.freeze_panes = "E2"
    sc.auto_filter.ref = f"A1:AA{last}"

    # Cable BOM
    bm = wb.create_sheet("Cable BOM")
    bcols = ["Product key", "Class", "Product (typical Southwire family)", "Conductors", "Size", "Cables", "Total length (ft)",
             "Waste", "Order length (ft)", "Standard reel (ft)", "Reels", "Order length (km)"]
    for c, h in enumerate(bcols, 1):
        x = bm.cell(1, c, h)
        x.font, x.fill = hdr_font, hdr_fill
        x.alignment = Alignment(wrap_text=True, vertical="center")
    keys = sorted({(r["prod"], r["cond"], r["size"]) for r in ROWS},
                  key=lambda k: (list(CLASS.values()).index(CLASS[k[0]]), k[0], SIZES_LV.index(k[2]) if k[2] in SIZES_LV else 99,
                                 k[1]))
    for i, (p, cnd, s) in enumerate(keys, 2):
        key = f"{p}|{cnd}|{s}"
        vals = [key, CLASS[p], P[p], cnd, AWG(s) if s != "-" else "-"]
        for c, v in enumerate(vals, 1):
            bm.cell(i, c, v).font = blk
        bm.cell(i, 6, f"=SUMIFS('Cable Schedule'!$AA$2:$AA${last},'Cable Schedule'!$K$2:$K${last},A{i})").font = grn
        bm.cell(i, 7, f"=SUMIFS('Cable Schedule'!$W$2:$W${last},'Cable Schedule'!$K$2:$K${last},A{i})").font = grn
        bm.cell(i, 8, "='Read Me'!$B$14").font = grn
        bm.cell(i, 9, f"=ROUNDUP(G{i}*(1+H{i}),-2)").font = blk
        x = bm.cell(i, 10, REEL[p])
        x.font = inp
        bm.cell(i, 11, f"=IF(J{i}>0,ROUNDUP(I{i}/J{i},0),0)").font = blk
        bm.cell(i, 12, f"=I{i}*0.3048/1000").font = blk
        for c in (7, 9, 10):
            bm.cell(i, c).number_format = "#,##0"
        bm.cell(i, 8).number_format = "0%"
        bm.cell(i, 12).number_format = "#,##0.0"
    n = len(keys) + 1
    bm.cell(n + 1, 3, "Total").font = Font(name=F, bold=True)
    for c, col in ((6, "F"), (7, "G"), (9, "I"), (11, "K"), (12, "L")):
        x = bm.cell(n + 1, c, f"=SUM({col}2:{col}{n})")
        x.font = Font(name=F, bold=True)
        x.number_format = "#,##0" if c != 12 else "#,##0.0"
    bm.cell(n + 3, 1, "Standard reel lengths are typical inputs (blue); edit to the actual reel / SIMpull reel sizes.").font = Font(
        name=F, italic=True, size=9)
    for c, w in enumerate([24, 15, 70, 10, 12, 7, 15, 7, 15, 12, 7, 12], 1):
        bm.column_dimensions[get_column_letter(c)].width = w
    bm.freeze_panes = "B2"

    # Accessories
    ac = wb.create_sheet("Accessories")
    rows = [("MV termination kits (outdoor / indoor, per conductor or per 3/C end)", f"=SUM('Cable Schedule'!X2:X{last})"),
            ("ARMOR-X / Type W connectors and glands", f"=SUM('Cable Schedule'!Y2:Y{last})"),
            ("MV splice kits (allowance: 1 per 2,500 ft of MV 1/C)",
             f"=ROUNDUP(SUMIFS('Cable Schedule'!W2:W{last},'Cable Schedule'!I2:I{last},\"MV*\",'Cable Schedule'!L2:L{last},\"1/C\")/2500,0)"),
            ("230 kV outdoor terminations", f"=2*SUMIFS('Cable Schedule'!O2:O{last},'Cable Schedule'!I2:I{last},\"HV 230 kV\")"),
            ("600 V compression lugs (power, 4 per end)",
             f"=2*4*SUMIFS('Cable Schedule'!N2:N{last},'Cable Schedule'!I2:I{last},\"LV 600 V\")"),
            ("Instrument / control cable glands (2 per run)",
             f"=2*(COUNTIF('Cable Schedule'!I2:I{last},\"Control\")+COUNTIF('Cable Schedule'!I2:I{last},\"Instrumentation\"))"),
            ("Fibre patch / termination (2 per link)", f"=2*COUNTIF('Cable Schedule'!I2:I{last},\"Fibre / data\")"),
            ("Ground connections: exothermic welds (allowance 1 per 25 ft of grid)",
             f"=ROUNDUP(SUMIFS('Cable Schedule'!W2:W{last},'Cable Schedule'!I2:I{last},\"Grounding\")/25,0)")]
    for c, h in enumerate(["Accessory", "Quantity"], 1):
        x = ac.cell(1, c, h)
        x.font, x.fill = hdr_font, hdr_fill
    for i, (t, f) in enumerate(rows, 2):
        ac.cell(i, 1, t).font = blk
        x = ac.cell(i, 2, f)
        x.font = grn
        x.number_format = "#,##0"
    ac.column_dimensions["A"].width = 80
    ac.column_dimensions["B"].width = 14

    # Summary by zone x class
    sm = wb.create_sheet("Summary")
    classes = list(dict.fromkeys(CLASS.values()))
    zones = sorted({(r["dst"]["area"] or "C") for r in ROWS})
    sm.cell(1, 1, "Total cable length (ft) by zone and class").font = Font(name=F, bold=True, size=11)
    sm.cell(2, 1, "Zone").font = hdr_font
    sm.cell(2, 1).fill = hdr_fill
    for c, k in enumerate(classes, 2):
        x = sm.cell(2, c, k)
        x.font, x.fill = hdr_font, hdr_fill
        x.alignment = Alignment(wrap_text=True)
    x = sm.cell(2, len(classes) + 2, "Total")
    x.font, x.fill = hdr_font, hdr_fill
    for r, z in enumerate(zones, 3):
        sm.cell(r, 1, f"{z}  {AREA.get(z, '')}").font = blk
        for c, k in enumerate(classes, 2):
            x = sm.cell(r, c, f"=SUMIFS('Cable Schedule'!$W$2:$W${last},'Cable Schedule'!$B$2:$B${last},\"{z}\","
                              f"'Cable Schedule'!$I$2:$I${last},\"{k}\")")
            x.font, x.number_format = grn, "#,##0"
        x = sm.cell(r, len(classes) + 2, f"=SUM(B{r}:{get_column_letter(len(classes) + 1)}{r})")
        x.font, x.number_format = Font(name=F, bold=True, size=9), "#,##0"
    tr = len(zones) + 3
    sm.cell(tr, 1, "Total").font = Font(name=F, bold=True)
    for c in range(2, len(classes) + 3):
        L_ = get_column_letter(c)
        x = sm.cell(tr, c, f"=SUM({L_}3:{L_}{tr - 1})")
        x.font, x.number_format = Font(name=F, bold=True, size=9), "#,##0"
    sm.cell(tr + 1, 1, "Total (miles)").font = Font(name=F, bold=True)
    x = sm.cell(tr + 1, len(classes) + 2, f"={get_column_letter(len(classes) + 2)}{tr}/5280")
    x.font, x.number_format = Font(name=F, bold=True, size=9), "#,##0.0"
    sm.cell(tr + 3, 1, "Schedule rows").font = Font(name=F, bold=True)
    sm.cell(tr + 3, 2, f"=COUNTIF('Cable Schedule'!B2:B{last},\"?\")").font = grn
    sm.cell(tr + 4, 1, "Cables (count)").font = Font(name=F, bold=True)
    sm.cell(tr + 4, 2, f"=SUM('Cable Schedule'!AA2:AA{last})").font = grn
    sm.cell(tr + 4, 2).number_format = "#,##0"
    sm.cell(tr + 6, 1, "Device count behind the schedule (typical for each equipment class)").font = Font(name=F, bold=True)
    for k_, (nm_, v_) in enumerate([("Analog instruments", N_DEV["ai"]), ("Discrete devices", N_DEV["di"]),
                                    ("Thermocouples", N_DEV["tc"]), ("RTDs", N_DEV["rtd"]), ("Solenoids", N_DEV["sol"]),
                                    ("LV motors", N_DEV["motors"]), ("Motor-operated valves", N_DEV["MOVs"])]):
        sm.cell(tr + 7 + k_, 1, nm_).font = blk
        sm.cell(tr + 7 + k_, 2, v_).font = inp
    sm.column_dimensions["A"].width = 34
    for c in range(2, len(classes) + 3):
        sm.column_dimensions[get_column_letter(c)].width = 13

    # Load list
    ll = wb.create_sheet("Load List")
    lcols = ["Load (to)", "Tag", "Zone", "Service", "Voltage (V)", "Load (kW)", "Current (A)", "Source (from)", "Product",
             "Size", "Sets"]
    for c, h in enumerate(lcols, 1):
        x = ll.cell(1, c, h)
        x.font, x.fill = hdr_font, hdr_fill
    pw = [r for r in ROWS if r["kw"]]
    for i, r in enumerate(pw, 2):
        vals = [r["dst"]["name"][:70], r["dst"].get("tag") or "", r["dst"]["area"], r["service"], r["volt"], r["kw"], r["amps"],
                r["src"]["name"][:60], CLASS[r["prod"]], AWG(r["size"]), r["sets"]]
        for c, v in enumerate(vals, 1):
            ll.cell(i, c, v).font = blk
    for c, w in enumerate([50, 12, 6, 46, 10, 10, 10, 44, 14, 12, 6], 1):
        ll.column_dimensions[get_column_letter(c)].width = w
    ll.freeze_panes = "A2"

    au = wb.create_sheet("Audit", 1)
    for c, h in enumerate(["Check", "Result", "Count", "Notes / examples"], 1):
        x = au.cell(1, c, h)
        x.font, x.fill = hdr_font, hdr_fill
    colr = {"PASS": "008000", "FAIL": "C00000", "CHECK": "B07000", "INFO": "404040"}
    for i, (a_, b_, c_, d_) in enumerate(audit(), 2):
        au.cell(i, 1, a_).font = blk
        au.cell(i, 2, b_).font = Font(name=F, bold=True, size=9, color=colr[b_])
        au.cell(i, 3, c_).font = blk
        au.cell(i, 4, d_).font = blk
        for c in (1, 4):
            au.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate([62, 9, 9, 110], 1):
        au.column_dimensions[get_column_letter(c)].width = w
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    export_xray()
    out = os.path.join(HERE, "brochure", "SK-3X1_Cable_Schedule_and_BOM.xlsx")
    wb.save(out)
    return out, last - 1, len(keys)


if __name__ == "__main__":
    out, n, k = write()
    print("wrote", out, n, "cable runs,", k, "BOM lines")
