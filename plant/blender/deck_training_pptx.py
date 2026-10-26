"""PowerPoint copy of the Nov 11 training deck: one full-bleed picture per slide (rendered from the PDF at 1920 x 1080),
with the run-of-show talking points for that slide as speaker notes. The PDF stays the master; rebuild this after
deck_training.py."""
import io
import os
import re

import pymupdf
from pptx import Presentation
from pptx.util import Emu

PLANT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF = os.path.join(PLANT, "renders", "deck", "training", "SW-3X1_Nov11_training.pdf")
ROS = os.path.join(PLANT, "presentation", "Nov11_run_of_show.md")
OUT = os.path.join(PLANT, "renders", "deck", "training", "SW-3X1_Nov11_training.pptx")


def notes_by_slide():
    """Talking points keyed by slide number, from the run of show's '**Slide N' and '**Slides a-b' entries."""
    notes = {}
    for line in open(ROS):
        m = re.match(r"^(?:- )?\*\*Slides? (\d+)(?:-(\d+))?[^*]*\*\*(.*)", line.strip())
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        head = re.match(r"^(?:- )?\*\*([^*]+)\*\*", line.strip()).group(1)
        text = re.sub(r"\*\*|`", "", head + " " + m.group(3)).strip()
        for k in range(a, b + 1):
            notes[k] = (notes.get(k, "") + "\n\n" + text).strip()
    return notes


def main():
    doc = pymupdf.open(PDF)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)       # 13.333 x 7.5 in, 16:9
    blank = prs.slide_layouts[6]
    notes = notes_by_slide()
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(1920 / page.rect.width, 1920 / page.rect.width))
        buf = io.BytesIO(pix.tobytes("jpeg", jpg_quality=88))
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        if i + 1 in notes:
            s.notes_slide.notes_text_frame.text = notes[i + 1]
    prs.save(OUT)
    print(f"wrote {OUT}: {len(doc)} slides, notes on {sum(1 for k in notes if k <= len(doc))}")


if __name__ == "__main__":
    main()
